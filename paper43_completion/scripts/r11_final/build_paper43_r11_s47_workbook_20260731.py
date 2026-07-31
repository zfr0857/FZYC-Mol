from copy import copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(r"D:\fzyc")
SRC = ROOT / "output" / "paper43_jcheminform_completion_20260726" / "JoC_R10_TECHNICAL_CANDIDATE_CONFLICT_HOLD_20260731" / "05_Additional_Files" / "Additional_file_2_Machine_readable_Supplementary_Tables_S1-S46_TAU_REVISED.xlsx"
OUT = ROOT / "output" / "paper43_jcheminform_completion_20260726" / "author_review" / "Additional_file_2_Machine_readable_Supplementary_Tables_S1-S47_R11.xlsx"

rows = [
    ("S_main", "number of primary split seeds (10)"),
    ("S_secondary", "number of seeds in explicitly labelled secondary analyses (5)"),
    ("s", "split seed"), ("F", "number of outer folds per seed (3)"), ("f", "outer fold"),
    ("u = (s, f)", "outer audit unit"), ("U_-s", "outer audit units from seeds other than held-out seed s"),
    ("N_-s", "number of units in U_-s; N_-s = |U_-s|"), ("j", "candidate"), ("K", "candidate count"),
    ("C_K", "eligible registered candidate prefix at size K"), ("V(u, j)", "inner-validation utility"),
    ("A(u, j)", "outer-audit utility (negative RMSE direction for regression)"),
    ("tau", "practical-equivalence reporting tolerance"),
    ("epsilon_num", "numerical-stability constant, fixed at 10^−12"),
    ("j_hat_u(K)", "validation-selected candidate at size K"),
    ("j_ref^inv(−s)", "K-invariant full-registry reference selected without seed s"),
    ("j_ref^dep(−s, K)", "K-dependent eligible-prefix reference selected without seed s"),
    ("G_inv(u, K)", "K-invariant full-registry completion gap"),
    ("G_dep(u, K)", "K-dependent within-prefix selection gap"),
    ("G_same(u, K)", "same-fold finite-set opportunity gap"),
    ("G_avail(u, K)", "availability component; G_inv = G_avail + G_dep"),
    ("G_bar_inv,e,s(K)", "seed-level mean of G_inv across the three outer folds"),
    ("Delta_inv(e)", "endpoint-level K=32 minus K=4 contrast averaged over seed blocks"),
    ("S_cov", "sample covariance matrix in Ledoit–Wolf shrinkage"),
    ("lambda_i", "non-negative shrinkage-correlation eigenvalue"),
    ("p_i", "unit-sum eigenvalue proportion"), ("pi_j", "selection proportion for candidate j"),
]

wb = load_workbook(SRC)
if "S47_Notation" in wb.sheetnames:
    del wb["S47_Notation"]
ws = wb.create_sheet("S47_Notation")
ws.append(["Symbol", "Definition"])
for row in rows: ws.append(row)
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:B{ws.max_row}"
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="315E8A"); cell.alignment = Alignment(horizontal="center")
ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 78
for row in ws.iter_rows(min_row=2):
    row[0].font = Font(name="Cambria Math"); row[1].font = Font(name="Times New Roman")
    row[1].alignment = Alignment(wrap_text=True, vertical="top")

idx = wb["Paper43_Index"]
for r in list(idx.iter_rows(min_row=2)):
    if r[0].value == "S47": idx.delete_rows(r[0].row)
idx.append(["S47", "S47_Notation", "Complete notation used in the main and supplementary analyses", "=COUNTA(S47_Notation!A:A)-1"])
for cell in idx[idx.max_row]:
    if idx.max_row > 2:
        source = idx[idx.max_row - 1][cell.column - 1]
        cell._style = copy(source._style); cell.number_format = source.number_format
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT, OUT.stat().st_size, ws.max_row - 1, idx.max_row)
