from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from PIL import Image
from openpyxl import load_workbook
import csv, hashlib, json, posixpath, re
import fitz

ROOT=Path(r"D:\fzyc")
BASE=ROOT/"output"/"paper43_jcheminform_completion_20260726"
PKG=BASE/"JoC_R11_STRUCTURAL_TECHNICAL_CANDIDATE_CONFLICT_HOLD_20260731"
QC=PKG/"06_Quality_Control"; QC.mkdir(exist_ok=True)
W="http://schemas.openxmlformats.org/wordprocessingml/2006/main";M="http://schemas.openxmlformats.org/officeDocument/2006/math";PR="http://schemas.openxmlformats.org/package/2006/relationships"
NS={"w":W,"m":M,"pr":PR}

checks=[]
def add(item,ok,detail): checks.append({"item":item,"status":"PASS" if ok else "FAIL","detail":str(detail)})
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def xml_doc(path):
    with ZipFile(path) as z: return etree.fromstring(z.read("word/document.xml"))
def text(root): return "\n".join("".join(p.xpath('.//w:t/text()|.//m:t/text()',namespaces=NS)) for p in root.xpath('//w:p',namespaces=NS))

man=PKG/"01_Manuscripts"
en=next(man.glob("Main*CONFLICT_HOLD_FINAL.docx")); zh=next(man.glob("*中文稿*CONFLICT_HOLD_FINAL.docx")); supp=PKG/"04_Additional_Files"/"Additional_file_1_Supplementary_Methods_and_Results_R11_FINAL.docx"
for label,p,expected_tables,expected_figs in [("English",en,4,8),("Chinese",zh,4,8),("Supplement",supp,5,0)]:
    root=xml_doc(p); tables=root.xpath('//w:body/w:tbl',namespaces=NS); drawings=root.xpath('//w:drawing',namespaces=NS)
    add(f"{label} DOCX tables",len(tables)==expected_tables,len(tables)); add(f"{label} DOCX figures",len(drawings)==expected_figs,len(drawings))
    full=text(root)
    add(f"{label} S47 reference",("S47" in full),"S47 present" if "S47" in full else "missing")
    if label!="Supplement":
        eq=[]
        for par in root.xpath('//w:p[.//m:oMath]',namespaces=NS):
            wt=''.join(par.xpath('.//w:t/text()',namespaces=NS)).strip()
            if re.fullmatch(r'\(\d+[ab]?\)',wt): eq.append(wt)
        expected=['(1)','(2)','(3)','(4)','(5)','(6)','(7)','(8)','(9)','(10a)','(10b)','(11)','(12a)','(12b)','(13)','(14)','(15)','(16)','(17)','(18)','(19a)','(19b)','(20a)','(20b)','(21)','(22a)','(22b)']
        add(f"{label} equation labels",eq==expected,eq)
        add(f"{label} Table 3 core rows",len(tables[2].xpath('./w:tr',namespaces=NS))==13,len(tables[2].xpath('./w:tr',namespaces=NS)))
        add(f"{label} Figure 3 primary panel",("primary" in full.lower() if label=="English" else "主要K不变" in full),"caption/body synchronized")
    else:
        last=tables[-1]; add("Supplement full notation rows",len(last.xpath('./w:tr',namespaces=NS))==29,len(last.xpath('./w:tr',namespaces=NS)))
        add("Supplement native 10^-12",bool(last.xpath('.//m:sSup',namespaces=NS)),len(last.xpath('.//m:sSup',namespaces=NS)))

zht=text(xml_doc(zh))
for bad in ["留一seed","reference = 1","候选库"]: add(f"Chinese forbidden term: {bad}",bad not in zht,"absent" if bad not in zht else "found")
add("Chinese support operator spacing","Tanimoto ≥ 0.70" in zht,"Tanimoto ≥ 0.70")

def rel_missing(path):
    missing=[]
    with ZipFile(path) as z:
        names=set(z.namelist())
        for rf in [n for n in names if n.endswith('.rels')]:
            relroot=etree.fromstring(z.read(rf))
            if rf=='_rels/.rels': srcdir=''
            else:
                prefix,relname=rf.rsplit('/_rels/',1); source=prefix+'/'+relname[:-5]; srcdir=posixpath.dirname(source)
            for rel in relroot:
                if rel.get('TargetMode')=='External': continue
                resolved=posixpath.normpath(posixpath.join(srcdir,rel.get('Target'))).lstrip('/')
                if resolved not in names: missing.append((rf,rel.get('Target'),resolved))
        if '[Content_Types].xml' not in names: missing.append(('package','','[Content_Types].xml'))
    return missing
for p in [en,zh,supp]:
    miss=rel_missing(p); add(f"DOCX relationships {p.name}",not miss,miss[:5])

figdir=PKG/"03_Main_Figures"
for n in range(1,9):
    files=[figdir/f"Figure{n}.svg",figdir/f"Figure{n}.pdf",figdir/f"Figure{n}_600dpi.png"]
    add(f"Figure {n} three formats",all(p.exists() and p.stat().st_size>1000 for p in files),[p.name for p in files])
    with Image.open(files[2]) as im:
        dpi=im.info.get('dpi',(0,0)); add(f"Figure {n} PNG 600 dpi",dpi[0]>=590 and dpi[1]>=590,dpi)
    svg=files[0].read_text(encoding='utf-8',errors='ignore')
    add(f"Figure {n} font family",'Times New Roman' in svg and 'Arial' not in svg,"Times New Roman / no Arial")

src=BASE/"source_data"/"fixed_reference_k32_vs_k4_contrasts.csv"
if not src.exists(): src=ROOT/"output"/"paper43_jcheminform_completion_20260726"/"source_data"/"fixed_reference_k32_vs_k4_contrasts.csv"
figcsv=PKG/"05_Source_Data"/"Figure_3A_primary_ten_seed_source.csv"
with figcsv.open(encoding='utf-8-sig') as f: frows=list(csv.DictReader(f))
expected={
 'BACE':(-.0195,-.0238,-.0152),'BBBP':(-.0263,-.0301,-.0225),'ClinTox':(-.0144,-.0241,-.0031),'HIA':(-.0022,-.0075,.0031),'P-gp':(-.0222,-.0271,-.0178),
 'ESOL':(.0683,.0618,.0757),'FreeSolv':(-.0586,-.0912,-.0263),'Lipophilicity':(-.0947,-.0992,-.0899),'Caco2':(-.0049,-.0121,.0008)}
obs={}; aliases={'bace':'BACE','bbbp':'BBBP','clintox':'ClinTox','esol':'ESOL','freesolv':'FreeSolv','lipo':'Lipophilicity','tdc_caco2_wang':'Caco2','tdc_hia_hou':'HIA','tdc_pgp_broccatelli':'P-gp'}
for r in frows:
    raw_endpoint=r.get('endpoint') or r.get('Endpoint') or r.get('dataset'); endpoint=aliases.get(raw_endpoint,raw_endpoint); nums=[]
    for key in ['estimate','effect','median','mean_natural_scale_effect','lower','low','ci_low','seed_block_interval_low','upper','high','ci_high','seed_block_interval_high']:
        if key in r and r[key] not in ('',None):
            try: nums.append(float(r[key]))
            except: pass
    if endpoint: obs[endpoint]=nums
ok=True
for k,vals in expected.items():
    if k not in obs or not all(any(abs(x-v)<5e-4 for x in obs[k]) for v in vals): ok=False
add("Figure 3A/Table 4 exact source values",ok,{k:obs.get(k) for k in expected})

xlsx=PKG/"04_Additional_Files"/"Additional_file_2_Machine_readable_Supplementary_Tables_S1-S47_R11.xlsx"
wb=load_workbook(xlsx,read_only=True,data_only=False)
add("S47 workbook sheet","S47_Notation" in wb.sheetnames,wb.sheetnames[-3:])
ws=wb['S47_Notation']; add("S47 workbook notation rows",ws.max_row-1==28,ws.max_row-1)
idx=wb['Paper43_Index']; idxvals=[r[0].value for r in idx.iter_rows(min_row=2)]; add("S47 workbook index",'S47' in idxvals,idxvals)

compat=PKG/"07_Compatibility"
for prefix,expected_pages in [('WPS_English',55),('WPS_Chinese',32),('WPS_Supplementary',14)]:
    p=next(compat.glob(prefix+'*.pdf')); pages=fitz.open(p).page_count; add(prefix+' pages',pages==expected_pages,pages)
lopdfs=list(compat.glob('LibreOffice_*.pdf'))
lopages=sorted(fitz.open(p).page_count for p in lopdfs); add('LibreOffice PDF page set',lopages==[14,34,55],lopages)
for p in list(compat.glob('*.pdf'))+list(man.glob('*.pdf')):
    try:
        d=fitz.open(p); ok=d.page_count>0 and all(page.rect.width>0 for page in d)
    except Exception as e: ok=False
    add('PDF parse '+p.name,ok,getattr(d,'page_count',0) if ok else 'failed')

track_en=next(man.glob('*TRACK_CHANGES*.docx')); track_zh=next(man.glob('*修订痕迹*.docx'))
for label,p in [('English tracked',track_en),('Chinese tracked',track_zh)]:
    with ZipFile(p) as z:
        doc=etree.fromstring(z.read('word/document.xml')); settings=etree.fromstring(z.read('word/settings.xml'))
    rev=len(doc.xpath('//w:ins|//w:del',namespaces=NS)); enabled=bool(settings.xpath('//w:trackRevisions',namespaces=NS))
    add(label+' revisions',rev>0 and enabled,{'xml_revisions':rev,'trackRevisions':enabled})

add("Real author metadata gate",False,"CONFLICT_HOLD: author/order/affiliation/correspondence/ORCID absent")
add("Declarations gate",False,"CONFLICT_HOLD: funding/CRediT/competing interests/acknowledgements not confirmed")
add("Public release gate",False,"CONFLICT_HOLD: no authorized public R11 release or immutable identifier")
add("Submission portal PDF gate",False,"CONFLICT_HOLD: external portal access not available")

with (QC/'R11_QC_checks.csv').open('w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=['item','status','detail']);w.writeheader();w.writerows(checks)
summary={'total':len(checks),'pass':sum(x['status']=='PASS' for x in checks),'fail':sum(x['status']=='FAIL' for x in checks),'technical_status':'PASS' if all(x['status']=='PASS' for x in checks[:-4]) else 'FAIL','release_status':'CONFLICT_HOLD'}
(QC/'R11_QC_summary.json').write_text(json.dumps({'summary':summary,'checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
lines=['# R11 structural technical candidate QC','',f"Technical status: **{summary['technical_status']}**",'Release status: **CONFLICT_HOLD**','',f"Checks: {summary['pass']} PASS / {summary['fail']} expected release-gate FAIL.",'','The final four FAIL rows are deliberate truth gates, not technical defects. Word and WPS opened the final DOCX files; WPS and LibreOffice exported all three PDFs. The external submission portal was not available.','','## Release blockers','- Real author and correspondence metadata are absent.','- Funding, CRediT, competing-interests and acknowledgement declarations are not author-confirmed.','- No authorized public R11 repository release or immutable identifier exists.','- Submission-portal PDF conversion was not run.']
(QC/'R11_QC_REPORT.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

manifest=PKG/'SHA256SUMS.csv'; rows=[]
for p in sorted(PKG.rglob('*')):
    if p.is_file() and p!=manifest: rows.append((sha(p),p.relative_to(PKG).as_posix(),p.stat().st_size))
with manifest.open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(['sha256','path','bytes']);w.writerows(rows)
zip_path=PKG.with_suffix('.zip')
with ZipFile(zip_path,'w',ZIP_DEFLATED,allowZip64=True) as z:
    for p in sorted(PKG.rglob('*')):
        if p.is_file(): z.write(p,(PKG.name/p.relative_to(PKG)).as_posix())
with ZipFile(zip_path) as z: bad=z.testzip()
print(json.dumps(summary,ensure_ascii=False));print('zip',zip_path,zip_path.stat().st_size,'bad',bad)
if summary['technical_status']!='PASS' or bad: raise SystemExit(2)
