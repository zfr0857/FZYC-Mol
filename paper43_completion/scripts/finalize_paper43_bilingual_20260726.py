from __future__ import annotations

import re
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"D:\fzyc")
PKG = ROOT / "output" / "paper43_jcheminform_completion_20260726"
SOURCE_EN = PKG / "author_review" / "Main_manuscript_Journal_of_Cheminformatics_CLEAN_PENDING_AUTHOR_METADATA_AND_REPOSITORY.docx"
OUT_EN = PKG / "author_review" / "Main_manuscript_Journal_of_Cheminformatics_FINAL_CLEAN.docx"
ZH_BASE_DIR = ROOT / "output" / "paper42_methodological_reframe_20260726"
OUT_ZH = PKG / "author_review" / "候选池扩张_有限验证下的机会与选择稳定性_中文同步终稿.docx"
FIG = PKG / "main_figures_submission"
CONTRASTS = PKG / "additional_files" / "tables" / "fixed_reference_k32_vs_k4_contrasts.csv"
DIVERSITY = PKG / "source_data" / "effective_diversity_10seed_summary.csv"


def set_east_asia(font, name: str) -> None:
    font.name = name
    font._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def get_style(doc: Document, name: str, base: str = "Normal"):
    if name not in doc.styles:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
    return doc.styles[name]


def configure_styles(doc: Document, chinese: bool = False) -> None:
    body_font = "宋体" if chinese else "Times New Roman"
    normal = doc.styles["Normal"]
    set_east_asia(normal.font, body_font); normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf.line_spacing = 2; pf.space_before = Pt(0); pf.space_after = Pt(0)
    title = doc.styles["Title"]
    set_east_asia(title.font, body_font); title.font.size = Pt(16); title.font.bold = True
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.line_spacing = 1
    for name, size, bold, italic in [("Heading 1", 14, True, False), ("Heading 2", 12, True, False), ("Heading 3", 12, True, True)]:
        if name in doc.styles:
            s = doc.styles[name]; set_east_asia(s.font, body_font); s.font.size = Pt(size); s.font.bold = bold; s.font.italic = italic
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; s.paragraph_format.keep_with_next = True; s.paragraph_format.keep_together = True
            s.paragraph_format.space_before = Pt(6); s.paragraph_format.space_after = Pt(3); s.paragraph_format.line_spacing = 1
    figcap = get_style(doc, "Figure Caption")
    set_east_asia(figcap.font, body_font); figcap.font.size = Pt(10); figcap.font.bold = False
    figcap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; figcap.paragraph_format.line_spacing = 1
    figcap.paragraph_format.space_before = Pt(6); figcap.paragraph_format.space_after = Pt(3); figcap.paragraph_format.keep_together = True
    tabcap = get_style(doc, "Table Caption")
    set_east_asia(tabcap.font, body_font); tabcap.font.size = Pt(10); tabcap.font.bold = True
    tabcap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; tabcap.paragraph_format.line_spacing = 1
    tabcap.paragraph_format.space_before = Pt(6); tabcap.paragraph_format.space_after = Pt(3); tabcap.paragraph_format.keep_with_next = True; tabcap.paragraph_format.keep_together = True
    tabtext = get_style(doc, "Table Text")
    set_east_asia(tabtext.font, body_font); tabtext.font.size = Pt(9)
    tabtext.paragraph_format.line_spacing = 1; tabtext.paragraph_format.space_before = Pt(0); tabtext.paragraph_format.space_after = Pt(0)
    tabnote = get_style(doc, "Table Note")
    set_east_asia(tabnote.font, body_font); tabnote.font.size = Pt(9)
    tabnote.paragraph_format.line_spacing = 1; tabnote.paragraph_format.space_before = Pt(2); tabnote.paragraph_format.space_after = Pt(6); tabnote.paragraph_format.keep_together = True
    if chinese:
        # Latin characters and numerals remain Times New Roman through the ASCII font slot.
        for s in [normal, title, figcap, tabcap, tabtext, tabnote, *[doc.styles[n] for n in ["Heading 1", "Heading 2", "Heading 3"] if n in doc.styles]]:
            s.font.name = "Times New Roman"
            s.font._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def clear_direct_font(run) -> None:
    rpr = run._r.rPr
    if rpr is None or run._r.xpath(".//m:oMath"): return
    for tag in ["w:rFonts", "w:sz", "w:szCs"]:
        for node in list(rpr.findall(qn(tag))): rpr.remove(node)


def remove_manual_breaks(doc: Document) -> None:
    for p in doc.paragraphs:
        ppr = p._p.get_or_add_pPr()
        for node in list(ppr.findall(qn("w:pageBreakBefore"))): ppr.remove(node)
        for br in list(p._p.xpath(".//w:br[@w:type='page']")): br.getparent().remove(br)


def style_paragraphs(doc: Document, chinese: bool = False) -> None:
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if i == 0:
            p.style = doc.styles["Title"]
        elif re.match(r"^[1-5](?:\.\d+)?\s+", text) or text in {"Abstract", "摘要", "Declarations", "声明", "References", "参考文献"}:
            level = 1 if re.match(r"^[1-5]\s+", text) or text in {"Abstract", "摘要", "Declarations", "声明", "References", "参考文献"} else 2
            p.style = doc.styles[f"Heading {level}"]
        elif text.startswith(("Figure ", "图 ", "图")) and re.match(r"^(Figure\s+\d+|图\s*\d+)", text):
            p.style = doc.styles["Figure Caption"]
        elif text.startswith(("Table ", "表 ", "表")) and re.match(r"^(Table\s+\d+|表\s*\d+)", text):
            p.style = doc.styles["Table Caption"]
        elif p.style.name not in ["Heading 1", "Heading 2", "Heading 3", "Title", "Figure Caption", "Table Caption", "Table Note"]:
            p.style = doc.styles["Normal"]
        for run in p.runs: clear_direct_font(run)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = text
    for p in cell.paragraphs:
        p.style = "Table Text"
        p.paragraph_format.keep_together = True
        for r in p.runs: r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def math_text(text: str):
    obj = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = text
    run.append(node); obj.append(run)
    return obj


def math_subscript(base: str, subscript: str):
    obj = OxmlElement("m:oMath")
    sub = OxmlElement("m:sSub")
    sub_pr = OxmlElement("m:sSubPr"); sub.append(sub_pr)
    base_node = OxmlElement("m:e")
    base_run = OxmlElement("m:r"); base_text = OxmlElement("m:t"); base_text.text = base
    base_run.append(base_text); base_node.append(base_run); sub.append(base_node)
    sub_node = OxmlElement("m:sub")
    sub_run = OxmlElement("m:r"); sub_text = OxmlElement("m:t"); sub_text.text = subscript
    sub_run.append(sub_text); sub_node.append(sub_run); sub.append(sub_node)
    obj.append(sub)
    return obj


def math_epsilon_constant():
    obj = OxmlElement("m:oMath")
    lead = OxmlElement("m:r"); lead_text = OxmlElement("m:t"); lead_text.text = "ε = "
    lead.append(lead_text); obj.append(lead)
    sup = OxmlElement("m:sSup"); sup.append(OxmlElement("m:sSupPr"))
    base = OxmlElement("m:e"); base_run = OxmlElement("m:r"); base_text = OxmlElement("m:t"); base_text.text = "10"
    base_run.append(base_text); base.append(base_run); sup.append(base)
    exponent = OxmlElement("m:sup"); exp_run = OxmlElement("m:r"); exp_text = OxmlElement("m:t"); exp_text.text = "−12"
    exp_run.append(exp_text); exponent.append(exp_run); sup.append(exponent)
    obj.append(sup)
    return obj


def apply_native_inline_math(doc: Document) -> None:
    factories = {"ε = 10⁻¹²": math_epsilon_constant, "0 log 0 := 0": lambda: math_text("0 log 0 := 0")}
    for paragraph in doc.paragraphs:
        original = paragraph.text
        if not any(token in original for token in factories):
            continue
        paragraph.clear()
        remaining = original
        while remaining:
            hits = [(remaining.find(token), token) for token in factories if token in remaining]
            if not hits:
                paragraph.add_run(remaining); break
            index, token = min(hits)
            if index:
                paragraph.add_run(remaining[:index])
            paragraph._p.append(factories[token]())
            remaining = remaining[index + len(token):]


def set_cell_math(cell, expression: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "Table Text"
    if "_" in expression:
        base, subscript = expression.split("_", 1)
        paragraph._p.append(math_subscript(base, subscript))
    else:
        paragraph._p.append(math_text(expression))
    paragraph.paragraph_format.keep_together = True
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_decimal_alignment(cell, position_twips: int = 1050) -> None:
    """Align the first decimal separator in a numeric cell using a Word decimal tab."""
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ppr = p._p.get_or_add_pPr()
        for old in list(ppr.findall(qn("w:tabs"))): ppr.remove(old)
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "decimal"); tab.set(qn("w:pos"), str(position_twips)); tabs.append(tab); ppr.append(tabs)
        tab_run = OxmlElement("w:r"); tab_run.append(OxmlElement("w:tab"))
        insert_at = 1 if p._p.find(qn("w:pPr")) is not None else 0
        p._p.insert(insert_at, tab_run)


def set_three_line_table(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tblpr.append(borders)
    for edge in ["top", "bottom", "left", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None: node = OxmlElement(f"w:{edge}"); borders.append(node)
        node.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        if edge in {"top", "bottom"}: node.set(qn("w:sz"), "8"); node.set(qn("w:color"), "000000")
    for row_i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); trpr.append(cant)
        if row_i == 0:
            rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trpr.append(rep)
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            # Remove all inherited cell borders before adding only the header rule.
            for old in list(tcpr.findall(qn("w:tcBorders"))): tcpr.remove(old)
            if row_i == 0:
                b = OxmlElement("w:tcBorders"); tcpr.append(b)
                bot = OxmlElement("w:bottom"); bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6"); bot.set(qn("w:color"), "000000"); b.append(bot)
            for shd in list(tcpr.findall(qn("w:shd"))): tcpr.remove(shd)
            for p in cell.paragraphs:
                p.style = "Table Text"
                for r in p.runs: clear_direct_font(r)


def find_caption(doc: Document, prefix: str):
    return next((p for p in doc.paragraphs if p.text.strip().startswith(prefix)), None)


def find_figure_caption(doc: Document, number: int):
    pattern = re.compile(rf"^(?:Figure\s+{number}\.|图\s*{number}(?:[.．]|\s))")
    return next((p for p in doc.paragraphs if pattern.match(p.text.strip())), None)


def replace_figure(doc: Document, number: int) -> None:
    caption = find_figure_caption(doc, number)
    if caption is None: return
    prev = caption._p.getprevious()
    if prev is None: return
    for child in list(prev):
        if child.tag != qn("w:pPr"): prev.remove(child)
    from docx.text.paragraph import Paragraph
    picture_p = Paragraph(prev, caption._parent)
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER; picture_p.paragraph_format.keep_with_next = True; picture_p.paragraph_format.keep_together = True
    picture_p.add_run().add_picture(str(FIG / f"Figure{number}_600dpi.png"), width=Inches(6.45))


def insert_paragraph_after(element, text: str, style=None, doc: Document | None = None):
    p = OxmlElement("w:p"); element.addnext(p)
    from docx.text.paragraph import Paragraph
    para = Paragraph(p, doc._body if doc is not None else element.getparent())
    if style is not None: para.style = style
    para.add_run(text)
    return para


def revise_tables(doc: Document, chinese: bool = False) -> None:
    # Table 1: synchronized caption, mathematical minus signs and centred n.
    t1 = doc.tables[0]
    cap1 = find_caption(doc, "Table 1.") or find_caption(doc, "表 1.") or find_caption(doc, "表1.")
    if cap1:
        cap1.text = "表1. 主要数据集与终点指标。" if chinese else "Table 1. Primary datasets and endpoint metrics."
        cap1.style = doc.styles["Table Caption"]
    for row_i, row in enumerate(t1.rows):
        for col_i, cell in enumerate(row.cells):
            value = re.sub(r"(?<!\d)-(?=\d)", "−", cell.text)
            if chinese:
                value = value.replace(" positive ", " 阳性 ").replace("Positive", "阳性")
            set_cell_text(cell, value, bold=row_i == 0)
            if col_i == 1:
                for p in cell.paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Table 2: only the current primary exposure and explicitly labelled secondary exposures.
    t2 = doc.tables[1]
    table2_rows = [
        ("Audit component", "Registry", "Evaluation design", "Recorded exposure"),
        ("Controlled prefix audit", "32 Morgan candidates", "9 endpoints; K = 4, 8, 16, 32; 10 seeds; 3 outer × 3 inner folds", "34 560 candidate fits; 8 130.16 candidate-fit seconds"),
        ("Calibration and resampling controls", "Stored 32-candidate results", "5 000 permutations; 100 resamples per mode/K/seed", "No additional model fitting"),
        ("Matched multiview audit (secondary)", "12 multiview candidates", "9 endpoints; 5 seeds; all C(12,3) registered subsets", "6 480 candidate fits"),
        ("Split-regime transfer audit (secondary)", "32 Morgan candidates", "3 endpoints; seeded scaffold and Tanimoto-component splits", "5 760 candidate fits"),
        ("Registry-composition intervention (secondary)", "Three locked candidate registries", "6 endpoints; K = 4, 8, 16, 32; composition and stability controls", "1 080 outer units; 64 616.35 downstream seconds"),
    ]
    if chinese:
        table2_rows = [
            ("审计组成", "候选池", "评价设计", "记录的计算暴露"),
            ("受控前缀主要审计", "32个Morgan候选", "9端点；K = 4、8、16、32；10种子；3外层×3内层折", "34 560次候选拟合；8 130.16候选拟合秒"),
            ("校准与重采样控制", "已存储的32候选结果", "5 000次置换；每模式/K/种子100次重采样", "无新增模型拟合"),
            ("匹配多表征审计（次要）", "12个多表征候选", "9端点；5种子；全部C(12,3)登记子集", "6 480次候选拟合"),
            ("切分机制迁移审计（次要）", "32个Morgan候选", "3端点；种子化骨架与Tanimoto分量切分", "5 760次候选拟合"),
            ("候选库组成干预（次要）", "三类锁定候选池", "6端点；K = 4、8、16、32；组成与稳定性控制", "1 080个外层单元；64 616.35下游秒"),
        ]
    while len(t2.rows) < len(table2_rows): t2.add_row()
    while len(t2.rows) > len(table2_rows): t2._tbl.remove(t2.rows[-1]._tr)
    for i, row_data in enumerate(table2_rows):
        for j, value in enumerate(row_data): set_cell_text(t2.cell(i, j), value, bold=i == 0)
    cap2 = find_caption(doc, "Table 2.") or find_caption(doc, "表 2.") or find_caption(doc, "表2.")
    if cap2:
        cap2.text = "表2. 审计组成与记录的计算暴露。" if chinese else "Table 2. Audit components and recorded computational exposure."
        cap2.style = doc.styles["Table Caption"]

    notation = [
        ("Symbol", "Definition"), ("S_main", "number of primary split seeds (10)"),
        ("S_secondary", "number of seeds in explicitly labelled secondary analyses (5)"), ("s", "split seed"),
        ("F", "number of outer folds per seed (3)"), ("f", "outer fold"), ("u = (s, f)", "outer audit unit"),
        ("j", "candidate"), ("K", "candidate count"),
        ("C_K", "eligible registered candidate prefix at size K"), ("V(u, j)", "inner-validation utility"),
        ("A(u, j)", "outer-audit utility"), ("j_ref,32,−s", "K-invariant full-registry reference selected without seed s"),
        ("j_ref,K,−s", "K-dependent eligible-prefix reference selected without seed s"),
        ("G_inv", "K-invariant full-registry completion gap"), ("G_dep", "K-dependent eligible-prefix gap"),
        ("G_same", "same-fold finite-set opportunity gap"), ("G_avail", "availability component, with G_inv = G_avail + G_dep"),
    ]
    if chinese:
        notation = [
            ("符号", "定义"), ("S_main", "主要分析的划分种子数（10）"),
            ("S_secondary", "明确标注的次要分析所用种子数（5）"), ("s", "划分种子"),
            ("F", "每个种子的外层折数（3）"), ("f", "外层折"), ("u = (s, f)", "外层审计单元"),
            ("j", "候选"), ("K", "候选数量"),
            ("C_K", "规模K下符合资格的预注册候选前缀"), ("V(u, j)", "内层验证效用"),
            ("A(u, j)", "外层审计效用"), ("j_ref,32,−s", "不使用种子s选出的K不变完整候选库参考"),
            ("j_ref,K,−s", "不使用种子s选出的K依赖合资格前缀参考"),
            ("G_inv", "K不变完整候选库完成差距"), ("G_dep", "K依赖合资格前缀差距"),
            ("G_same", "同折有限集合机会差距"), ("G_avail", "可用性分量，满足G_inv = G_avail + G_dep"),
        ]
    t3 = doc.tables[2]
    while len(t3.rows) < len(notation): t3.add_row()
    while len(t3.rows) > len(notation): t3._tbl.remove(t3.rows[-1]._tr)
    for i, (a, b) in enumerate(notation):
        if i == 0:
            set_cell_text(t3.cell(i, 0), a, bold=True)
        else:
            set_cell_math(t3.cell(i, 0), a)
        set_cell_text(t3.cell(i, 1), b, bold=i == 0)
    cap3 = find_caption(doc, "Table 3.") or find_caption(doc, "表 3.") or find_caption(doc, "表3.")
    if cap3:
        cap3.text = "表3. 审计估计量使用的数学符号。" if chinese else "Table 3. Mathematical notation used in the audit estimands."
        cap3.style = doc.styles["Table Caption"]
        cap3.paragraph_format.page_break_before = True

    effects = pd.read_csv(CONTRASTS)
    effects = effects[effects.estimand.eq("fixed_k32_gap")].set_index("dataset")
    rows = [
        ("Endpoint", "K = 32 minus K = 4 contrast (95% sensitivity interval)", "Interpretation"),
        ("Classification: ROC-AUC completion-gap contrast", "", ""),
    ]
    for key, label in [("bace", "BACE"), ("bbbp", "BBBP"), ("clintox", "ClinTox"), ("tdc_hia_hou", "HIA"), ("tdc_pgp_broccatelli", "P-gp")]:
        r = effects.loc[key]; v, lo, hi = r.mean_natural_scale_effect, r.seed_block_interval_low, r.seed_block_interval_high
        interp = "Smaller gap at K = 32; interval excludes 0" if hi < 0 else "Smaller gap at K = 4; interval excludes 0" if lo > 0 else "Interval includes 0"
        rows.append((label, f"{v:.4f} ({lo:.4f}, {hi:.4f})".replace("-", "−"), interp))
    rows.append(("Regression: RMSE completion-gap contrast", "", ""))
    for key, label in [("esol", "ESOL"), ("freesolv", "FreeSolv"), ("lipo", "Lipophilicity"), ("tdc_caco2_wang", "Caco2")]:
        r = effects.loc[key]; v, lo, hi = r.mean_natural_scale_effect, r.seed_block_interval_low, r.seed_block_interval_high
        interp = "Smaller gap at K = 32; interval excludes 0" if hi < 0 else "Smaller gap at K = 4; interval excludes 0" if lo > 0 else "Interval includes 0"
        rows.append((label, f"{v:.4f} ({lo:.4f}, {hi:.4f})".replace("-", "−"), interp))
    if chinese:
        translated = [("端点", "K = 32减K = 4对比（95%敏感性区间）", "解释")]
        for endpoint, value, interp in rows[1:]:
            if endpoint.startswith("Classification:"):
                translated.append(("分类：ROC-AUC完成差距对比", "", "")); continue
            if endpoint.startswith("Regression:"):
                translated.append(("回归：RMSE完成差距对比", "", "")); continue
            interp = interp.replace("Smaller gap at K = 32; interval excludes 0", "K = 32下完成差距更小；区间不含0")
            interp = interp.replace("Smaller gap at K = 4; interval excludes 0", "K = 4下完成差距更小；区间不含0").replace("Interval includes 0", "区间包含0")
            translated.append((endpoint, value, interp))
        rows = translated
    t4 = doc.tables[3]
    for i, row_data in enumerate(rows):
        for j, value in enumerate(row_data):
            set_cell_text(t4.cell(i, j), value, bold=i == 0)
            if j == 1 and i not in {0, 1, 7}: set_decimal_alignment(t4.cell(i, j))
        if i in {1, 7}:
            merged = t4.cell(i, 0).merge(t4.cell(i, 2)); set_cell_text(merged, row_data[0], bold=True)
    caption = find_caption(doc, "Table 4.") or find_caption(doc, "表 4.") or find_caption(doc, "表4.")
    if caption:
        caption.text = ("表4. K不变完整候选库完成差距对比。" if chinese else "Table 4. K-invariant full-registry completion-gap contrasts.")
        caption.style = doc.styles["Table Caption"]
        caption.paragraph_format.page_break_before = True
        note_text = ("注：负值表示K = 32下相对于完整候选库参考的完成差距更小。区间为描述性划分种子敏感性区间。" if chinese else "Note: Negative values indicate a smaller completion gap relative to the full-registry reference at K = 32. Intervals are descriptive split-seed sensitivity intervals.")
        # Avoid duplicate note on repeated runs.
        nxt = t4._tbl.getnext()
        if nxt is None or note_text not in "".join(nxt.itertext()): insert_paragraph_after(t4._tbl, note_text, doc.styles["Table Note"], doc)

    for table in doc.tables: set_three_line_table(table)


def keep_figures_with_captions(doc: Document) -> None:
    for n in range(1, 9):
        cap = find_figure_caption(doc, n)
        if cap is None: continue
        cap.style = doc.styles["Figure Caption"]; cap.paragraph_format.keep_together = True
        prev = cap._p.getprevious()
        if prev is not None:
            ppr = prev.get_or_add_pPr(); keep = ppr.find(qn("w:keepNext"))
            if keep is None: ppr.append(OxmlElement("w:keepNext"))


def add_supplement_statement(doc: Document, chinese: bool = False) -> None:
    needle = "Additional file" if not chinese else "补充文件"
    if any("S1–S46" in p.text and "S1–S25" in p.text for p in doc.paragraphs): return
    target = next((p for p in doc.paragraphs if p.text.strip() in {"Declarations", "声明"}), None)
    if target is None: return
    text = ("补充材料在正文中对应说明如下：补充文件1含补充方法、补充结果和补充表说明；补充文件2含机器可读补充表S1–S46；补充文件3含补充图S1–S25；补充文件4含代码、环境、清单和复现说明。正文涉及的敏感性分析、构造对照、实用等价、化学支持及负结果均在相应补充表或补充图中给出可核查细节。" if chinese else
            "The supplementary package is cross-described here for transparency: Additional file 1 contains Supplementary Methods, Results and table notes; Additional file 2 contains machine-readable Tables S1–S46; Additional file 3 contains Figures S1–S25; and Additional file 4 contains code, environments, manifests and reproduction instructions. Sensitivity analyses, constructed controls, practical-equivalence results, chemical-support analyses and negative results mentioned in the main text are documented in the corresponding supplementary tables or figures.")
    p = OxmlElement("w:p"); target._p.addprevious(p)
    from docx.text.paragraph import Paragraph
    para = Paragraph(p, target._parent); para.style = doc.styles["Normal"]; para.add_run(text)


def synchronize_english_text(doc: Document) -> None:
    effects = pd.read_csv(CONTRASTS)
    effects = effects[effects.estimand.eq("fixed_k32_gap")]
    negative = int((effects.mean_natural_scale_effect < 0).sum())
    positive = int((effects.mean_natural_scale_effect > 0).sum())
    excludes = int(((effects.seed_block_interval_high < 0) | (effects.seed_block_interval_low > 0)).sum())
    diversity = pd.read_csv(DIVERSITY)
    d32 = diversity[(diversity.matrix_level.eq("outer")) & diversity.candidate_count.eq(32)].set_index("transformation")
    div = {key: float(d32.loc[key, "ledoit_wolf_entropy_rank_median"]) for key in
           ["raw", "row_centred", "fixed_reference_relative", "within_unit_rank"]}
    replace_start(doc, "Results:",
        f"Results: Against the K-invariant full-registry reference, K = 32 minus K = 4 completion-gap contrasts were negative in {negative} of 9 endpoints and positive in {positive}; negative contrasts mean a smaller full-registry completion gap at K = 32, whereas a positive contrast means a smaller gap at K = 4. Descriptive split-seed sensitivity intervals excluded zero for {excludes} endpoints. At K = 32, 79.6% of selections were ε-successes at the retrospectively locked middle reporting tolerance. PR-AUC and minority-constrained selectors changed the ROC-AUC-selected candidate in 52.7% and 66.0% of classification audit units, respectively; the minority-recall rule did not reliably attain the 0.80 target on held-out folds.")
    replace_start(doc, "The evaluation procedure for each candidate was held constant across K.",
        "The evaluation procedure for each candidate was held constant across K. The unique current primary ten-seed matrix comprised 34 560 candidate fits across seeds 11, 23, 37, 53, 71, 83, 97, 113, 127 and 149. The corresponding recorded current-run candidate-fit time was 8 130.16 seconds. Historical five-seed fits and timings are retained only in the supplementary provenance audit and are not added to the current primary exposure.")
    replace_start(doc, "Equations (1)",
        "Equations (1)–(4) define candidate selection, finite-audit selection loss and chance-adjusted ranking. Equations (5)–(6) define the cross-fitted references and gaps; Equations (7)–(10) define matrix transformations and effective ranks; Equations (11)–(13) define composition contrasts and normalized gains; and Equation (14) defines normalized selection entropy.")
    replace_start(doc, "We used ε =",
        "We used the numerical-stability constant ε = 10⁻¹². Ledoit–Wolf shrinkage used the scaled-identity target and the analytic coefficient implemented in scikit-learn. The shrinkage covariance was rescaled to a correlation matrix R before eigenanalysis. Correlation eigenvalues below zero only because of floating-point error were clipped to zero, with the convention 0 log 0 := 0. Paired normalized gains and cross-fitted gaps were computed within paired outer units, averaged over the three outer folds within each seed, and then summarized within endpoint using seed as the block; units with an absolute denominator at or below ε were reported as missing.")
    replace_start(doc, "For the ten-seed primary audit, each endpoint and K produced",
        "For the ten-seed primary effective-diversity audit, each endpoint and K produced a 30 × K outer-utility matrix and a 90 × K inner-utility matrix from ten split seeds, three outer folds and three inner folds per outer-training partition. Four prespecified transformations were analysed: raw utilities; row-centred utilities; fixed-reference-relative utilities; and within-unit ranks. The ten-seed outer matrix is shown in Figure 2 and Table S6; the corresponding inner matrix is reported in Table S6. Five-seed composition and multiview matrices remain separate secondary analyses and are not mixed with these estimates.")
    replace_start(doc, "Effective diversity was derived from candidate-utility matrices",
        "Effective diversity was derived from candidate-utility matrices rather than candidate labels or model-family counts. Outer matrices characterize behavioural similarity on held-out scaffold groups, whereas inner matrices assess whether redundancy is already visible during selection. Candidate columns with numerical variance at or below 10⁻¹² were excluded before correlation estimation and their post-transformation column counts were recorded. Fixed-reference-relative matrices subtract candidate 1 from every other candidate and therefore contain K − 1 columns.")
    replace_start(doc, "Uncertainty in endpoint-specific effective diversity",
        "Sensitivity of endpoint-specific effective diversity was assessed by omitting each of the ten split seeds, each outer-fold label and, for inner matrices, each inner-fold label, followed by complete recomputation of shrinkage correlations and effective ranks. Fixed-reference-relative estimates were also repeated with three additional registry-prespecified reference candidates; no reference was selected from outer performance. Across-endpoint summaries are medians, IQRs and ranges, without a population-level confidence interval because the nine endpoints were not treated as a random sample from a molecular-task population (Tables S6–S7).")
    replace_start(doc, "The five classification endpoints were refitted",
        "The five classification endpoints were refitted on the frozen ten-seed nested splits to record inner ROC-AUC, PR-AUC and operating-point metrics. In each inner fold, the minority label was the less frequent class in that fold's training partition. All distinct ROC thresholds from the inner validation predictions were evaluated. If one or more thresholds attained minority recall ≥0.80, the threshold maximizing majority recall was selected; ties were resolved by higher minority recall, higher balanced recall and then proximity to 0.5. If none attained 0.80, the same lexicographic rule was applied to all finite-recall thresholds and the failure of threshold feasibility was recorded. Candidate-specific thresholds were aggregated as the median of the three inner-fold thresholds, and the minority label as their mode, without post-hoc probability calibration. Candidate selection first retained candidates whose mean inner minority recall was ≥0.80; if none qualified, the full eligible prefix was retained. It then maximized mean inner majority recall, with ties resolved by mean inner minority recall, mean inner PR-AUC and registry order. The candidate identity, aggregated threshold and minority label were frozen before one-time outer-fold evaluation. We report threshold feasibility, candidate eligibility, outer target satisfaction, minority and majority recall, minority miss rate, PR-AUC and ROC-AUC.")
    replace_start(doc, "In the five-seed secondary effective-diversity reconstruction",
        f"In the ten-seed primary outer effective-diversity audit, the K = 32 endpoint medians of Ledoit–Wolf entropy rank were {div['raw']:.2f} for raw utilities, {div['row_centred']:.2f} after row centring, {div['fixed_reference_relative']:.2f} for fixed-reference-relative utilities and {div['within_unit_rank']:.2f} for within-unit ranks. These matrices retain common audit-unit difficulty, remove common level shifts, express contrasts to a prespecified reference or preserve ordering without utility spacing, respectively. The corresponding 30 × 32 outer and 90 × 32 inner matrices are reported without combining five-seed secondary composition analyses (Figure 2; Tables S6–S7).")
    replace_start(doc, "Estimator, seed, fold and reference sensitivities",
        "Matrix level, estimator, leave-one-seed, leave-one-fold and reference sensitivities changed magnitudes but preserved the distinction between nominal K and matrix-dependent effective rank. Full IQRs, ranges, participation-ratio ranks, shrinkage coefficients, correlation summaries and inner-matrix results are in Table S6; leave-one-unit and predefined-reference results are in Table S7.")
    replace_start(doc, "Holding the K = 32 cross-fitted reference identity fixed across K",
        "Holding the K = 32 cross-fitted reference identity fixed across K removed comparator drift. At K = 32, classification ε-success was 62.7%, 80.0% and 86.0% across the 0.005/0.010/0.020 ROC-AUC grid; regression ε-success was 73.3%, 79.2% and 82.5% across the 0.025/0.050/0.100 RMSE grid. The pooled task-appropriate working-middle-threshold success rate was 79.6%; the mean cross-fitted set size was 10.17, the selected-in-set rate was 85.9% and mean validation–cross-fit set Jaccard overlap was 0.490. These are retrospectively locked reporting tolerances, not clinical or pharmacological thresholds. At K = 32, PR-AUC selection changed the ROC-AUC-selected candidate in 52.7% of units and changed outer PR-AUC, ROC-AUC, minority recall, majority recall and minority miss rate by +0.0018, −0.0026, −0.0002, −0.0136 and +0.0002, respectively. The minority-recall-constrained selector changed candidates in 66.0% of units and changed the same outcomes by +0.0007, −0.0008, −0.0094, +0.0059 and +0.0094. An inner-eligible candidate existed in 100.0% of K = 32 units, whereas the frozen rule achieved outer minority recall ≥0.80 in 62.0%; the negative mean recall change was retained rather than interpreted as automatic constraint transfer (Tables S38–S40).")
    replace_start(doc, "Figure 2.",
        "Figure 2. Ten-seed matrix-dependent candidate diversity. (A) Ledoit–Wolf entropy rank across raw, row-centred, fixed-reference-relative and within-unit-rank outer-utility matrices; lines are endpoint medians, ribbons are endpoint IQRs and the grey dotted line is nominal K. (B) Entropy and participation-ratio concordance at K = 32 with a 45° reference. (C) Median candidate correlation after adjustment. (D) Fixed-reference-relative endpoint estimates at K = 32; solid and dashed ranges omit one seed and one outer-fold label, respectively, filled teal points are complete ten-seed estimates and hollow diamonds use a predefined linear reference. Each outer matrix contains 30 seed–fold rows; corresponding 90-row inner-matrix results are reported in Tables S6–S7.")
    replace_start(doc, "Figure 6.",
        "Figure 6. Prediction reliability across chemical-support boundaries. (A) Prediction correlations and high-error Jaccard overlaps are combined in a symmetric four-model matrix with separate in-panel colour strips. (B) Classification ROC-AUC, regression RMSE, classification false-negative rate and classification/regression high-error enrichment are consolidated into one Tanimoto-support risk matrix; cell text gives natural-scale medians and colour encodes only the within-row adverse direction. (C) Novel-scaffold relative changes in error overlap, model disagreement, high-error enrichment and false-negative enrichment are shown on a log-ratio axis. (D) Discrimination and minority-safety measures are separated for four fixed-configuration models; the purple dashed line is the 0.90 coverage target, and prediction-set size is mapped from its 1–2 scale only for display.")
    replace_start(doc, "Figure 7.",
        "Figure 7. Expanded candidate-pool composition intervention. (A) At K = 32, normalized gain uses paired homogeneous-audit-best normalization: the validation-selected gain is divided by the corresponding homogeneous-pool finite-audit-best gain; selected and finite-audit-best opportunity values are shown for six prespecified endpoints. (B) Validation-selected gain and cross-fitted gap are shown at K = 4, 8, 16 and 32 in separate classification and regression task bands; the pool legend is shown once above panel B. (C) CAHit@3 is shown for each endpoint, pool and K combination, with negative values retained; the rightmost Entropy column in the main figure reports normalized candidate-selection entropy at K = 32. H denotes homogeneous Morgan, MV classical multiview and M modern-augmented. (D) Normalized selected gain and downstream fitting/prediction time are compared under equal-K circles and equal-downstream-budget diamonds; marker area denotes K and the empirical Pareto frontier is shown. This composition intervention is a five-seed secondary analysis after averaging three outer folds within seed; time excludes model acquisition, encoder pretraining and cached embedding extraction.")
    replace_start(doc, "Against the K-invariant K = 32 leave-one-seed-out reference",
        f"Against the K-invariant K = 32 leave-one-seed-out reference, K = 32 minus K = 4 full-registry completion-gap contrasts were negative in {negative} of 9 endpoints and positive in {positive}. A negative contrast means that the full-registry completion gap was smaller at K = 32; a positive contrast means that it was smaller at K = 4. Descriptive split-seed sensitivity intervals excluded zero for {excludes} endpoints. Figure 3C and Table 4 are generated from the same endpoint source table and retain the same estimates, interval limits and signs without pooling classification and regression scales.")
    replace_start(doc, "Figure 8.",
        "Figure 8. Practical equivalence and metric-dependent selection in the ten-seed primary audit. (A) Classification ROC-AUC and regression RMSE gaps are shown on separate stacked axes sharing candidate count K; circles/solid lines, squares/dashed lines and triangles/dotted lines denote K-invariant, K-dependent and same-fold estimands, respectively, so estimands remain distinguishable in greyscale. (B) Cross-fitted ε-success is shown across K at the retrospectively locked middle reporting tolerances. (C) Mean cross-fitted near-equivalent set size is shown on one axis; the classification/regression legend in panel B also applies to panel C. (D) Candidate switching and outer changes are shown for the five classification endpoints, with PR-AUC switch, recall-rule switch, ΔPR-AUC and Δminority recall labelled explicitly; negative minority-recall changes are retained. Panels A–D are arranged in a 2 × 2 layout.")
    replace_start(doc, "Additional file 4 (ZIP; .zip).",
        "Additional file 4 (ZIP; .zip). Title: Code and reproducibility package. Description: Source code, locked configurations, compact source results, manifests, equation-to-code mapping, integrity hashes and rerun entry points. Its portable base corresponds to release paper-release-2026-07-r9 and commit 9635a902fa3cc7bb7b71a234c1b2bbbe415193f0; the paper43 completion overlay is post-release and must be mirrored to a synchronized public release before submission.")
    replace_start(doc, "The datasets supporting this article are public",
        "The datasets supporting this article are public and the processed audit tables are included in Additional files 1–4. The latest verified public code release is https://github.com/zfr0857/FZYC-Mol/releases/tag/paper-release-2026-07-r9 at commit 9635a902fa3cc7bb7b71a234c1b2bbbe415193f0. Additional file 4 contains that portable base plus the post-release paper43 ten-seed completion overlay. Because the overlay is not yet present in the public tag, repository synchronization and an immutable archive remain submission blockers; no archive DOI is claimed in this version.")
    replace_start(doc, "Software record.",
        "Software record. Project name: FZYC-Mol candidate-pool audit. Project home page: https://github.com/zfr0857/FZYC-Mol. Archived version: paper-release-2026-07-r9, commit 9635a902fa3cc7bb7b71a234c1b2bbbe415193f0. Operating system: platform independent where the locked Python environment is supported. Programming language: Python 3.13.7. Licence: MIT. Restrictions on reuse: none beyond the licence and the source-dataset terms.")


def format_english() -> None:
    doc = Document(SOURCE_EN)
    configure_styles(doc); remove_manual_breaks(doc)
    synchronize_english_text(doc)
    for n in range(1, 9): replace_figure(doc, n)
    revise_tables(doc); apply_native_inline_math(doc); add_supplement_statement(doc); style_paragraphs(doc); keep_figures_with_captions(doc)
    # Keep Table 3 and its caption together when pagination permits.
    cap3 = find_caption(doc, "Table 3.")
    if cap3: cap3.paragraph_format.keep_with_next = True
    doc.core_properties.title = doc.paragraphs[0].text
    doc.save(OUT_EN)


def replace_start(doc: Document, prefix: str, text: str) -> bool:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): p.text = text; return True
    return False


def insert_before(target, heading: str, body: str, doc: Document) -> None:
    for text, style in [(heading, "Heading 2"), (body, "Normal")]:
        p = OxmlElement("w:p"); target._p.addprevious(p)
        from docx.text.paragraph import Paragraph
        para = Paragraph(p, target._parent); para.style = doc.styles[style]; para.add_run(text)


def build_chinese() -> None:
    candidates = [p for p in ZH_BASE_DIR.glob("*.docx") if any(ord(c) > 127 for c in p.name) and p.stat().st_size > 1_000_000]
    if not candidates: raise FileNotFoundError("Chinese base manuscript not found")
    doc = Document(candidates[0])
    diversity = pd.read_csv(DIVERSITY)
    d32 = diversity[(diversity.matrix_level.eq("outer")) & diversity.candidate_count.eq(32)].set_index("transformation")
    div = {key: float(d32.loc[key, "ledoit_wolf_entropy_rank_median"]) for key in
           ["raw", "row_centred", "fixed_reference_relative", "within_unit_rank"]}
    configure_styles(doc, chinese=True); remove_manual_breaks(doc)
    doc.paragraphs[0].text = "分子性质预测中有限验证下的候选池机会与选择稳定性：重复嵌套审计"
    # The prior Chinese working draft omitted the author block; restore transparent placeholders.
    anchor = doc.paragraphs[0]
    for value in ["[作者姓名与ORCID须在投稿前补充]", "[机构、完整邮寄地址须在投稿前补充]", "*通讯作者：[姓名、电子邮箱及邮寄地址须补充]", "短标题：有限验证下的候选池机会"]:
        anchor = insert_paragraph_after(anchor._p, value, doc.styles["Normal"], doc)
    replace_start(doc, "[AUTHOR", "[作者姓名与ORCID须在投稿前补充]")
    replace_start(doc, "Running title:", "短标题：有限验证下的候选池机会")
    replace_start(doc, "背景：", "背景：分子性质预测研究常比较不断扩大的、彼此相关的分子表征、学习器与调参变体候选库。扩张可能增加互补的化学信息，但有限验证数据未必能够稳定地排序这些机会。")
    replace_start(doc, "方法：", "方法：我们在9个公开端点、32个预注册轻量候选、K = 4、8、16和32、3个内层与3个外层骨架折以及10个划分种子下开展审计。主要分析采用K不变的K = 32留一划分种子交叉拟合参考；K依赖交叉拟合参考和同折最大值分别作为敏感性量与描述性机会量。另评估实用等价网格、PR-AUC与训练期少数类召回约束选择、重复/弱/互补构造对照、经验相关异方差恢复模拟、候选库组成、计算暴露及化学支持。")
    replace_start(doc, "结果：", "结果：相对于K不变完整候选库参考，9个端点中8个端点的K = 32减K = 4完成差距对比为负，1个为正；7个端点的描述性划分种子敏感性区间不含0。K = 32时，任务匹配中间容差下的交叉拟合ε成功率为79.6%。PR-AUC与少数类召回约束规则分别在52.7%和66.0%的分类审计单元中改变ROC-AUC所选候选；后者使外层少数类召回平均变化−0.0094，且未稳定达到0.80目标。")
    replace_start(doc, "结论：", "结论：候选池扩张同时改变了可获得的机会，以及有限验证能否稳定实现这些机会。结果取决于端点、候选库组成、选择指标、候选相关性、化学支持和划分机制。K不变交叉拟合、实用等价报告与指标匹配选择降低了解释歧义，但不能替代外部验证，也不能证明更大候选库具有普遍效应。")
    replace_start(doc, "科学贡献：", "科学贡献：本研究将K不变交叉拟合参考与K依赖参考和同折机会量明确分离，并量化分子基准中的实用等价与指标依赖选择。构造依赖对照和经验相关异方差模拟说明，名义候选数量不能替代有效搜索多样性。所提供的审计与报告包面向第三方复现，而非模型排行榜。")
    replace_start(doc, "关键词：", "关键词：分子性质预测；候选池扩张；嵌套交叉验证；交叉拟合参考；实用等价；PR-AUC；选择稳定性；化学支持")
    replace_start(doc, "Background:", "背景：分子性质预测研究常比较不断扩大的、彼此相关的分子表征、学习器与调参变体候选库。扩张可能增加互补的化学信息，但有限验证数据未必能够稳定地排序这些机会。")
    replace_start(doc, "Methods:", "方法：我们在9个公开端点、32个预注册轻量候选、K = 4、8、16和32、3个内层与3个外层骨架折以及10个划分种子下开展审计。主要分析采用K不变的K = 32留一划分种子交叉拟合参考；K依赖交叉拟合参考和同折最大值分别作为敏感性量与描述性机会量。另评估实用等价网格、PR-AUC与训练期少数类召回约束选择、重复/弱/互补构造对照、经验相关异方差恢复模拟、候选库组成、计算暴露及化学支持。")
    replace_start(doc, "Results:", "结果：相对于K不变完整候选库参考，9个端点中8个端点的K = 32减K = 4完成差距对比为负，1个为正；7个端点的描述性划分种子敏感性区间不含0。K = 32时，任务匹配中间容差下的交叉拟合ε成功率为79.6%。PR-AUC与少数类召回约束规则分别在52.7%和66.0%的分类审计单元中改变ROC-AUC所选候选；后者使外层少数类召回平均变化−0.0094，且未稳定达到0.80目标。")
    replace_start(doc, "Conclusions:", "结论：候选池扩张同时改变了可获得的机会，以及有限验证能否稳定实现这些机会。结果取决于端点、候选库组成、选择指标、候选相关性、化学支持和划分机制。K不变交叉拟合、实用等价报告与指标匹配选择降低了解释歧义，但不能替代外部验证，也不能证明更大候选库具有普遍效应。")
    replace_start(doc, "Scientific Contribution:", "科学贡献：本研究将K不变交叉拟合参考与K依赖参考和同折机会量明确分离，并量化分子基准中的实用等价与指标依赖选择。构造依赖对照和经验相关异方差模拟说明，名义候选数量不能替代有效搜索多样性。所提供的审计与报告包面向第三方复现，而非模型排行榜。")
    replace_start(doc, "Keywords:", "关键词：分子性质预测；候选池扩张；嵌套交叉验证；交叉拟合参考；实用等价；PR-AUC；选择稳定性；化学支持")
    replace_start(doc, "2.8 ", "2.8 留一划分种子参考与机会差距分解")
    replace_start(doc, "因此，本研究开展回顾性锁定", "因此，本研究开展回顾性锁定、任务分层的审计，而非模型排行榜。核心问题是候选池扩张创造的机会能否由有限验证信息稳定兑现。主要分析采用九个端点、32个候选、K = 4、8、16、32和10个划分种子的K不变K = 32留一种子参考；多表征、候选库组成、可靠性及适用的切分迁移结果保留为明确标注的五种子次要分析。所有结论均限定于已评估的端点、候选池和切分机制。")
    replace_start(doc, "候选在不同K下沿用相同预处理", "候选在不同K下沿用相同预处理、内层折和重拟合逻辑，但候选池扩张增加总搜索暴露。当前主要十种子矩阵共34 560次候选拟合，覆盖种子11、23、37、53、71、83、97、113、127和149，记录候选拟合时间8 130.16秒。历史五种子拟合数与时间仅保留于补充数据溯源表，不计入当前主要分析。匹配多表征、切分迁移和候选库组成干预均在表2中明确标注为次要分析；完整资源见补充表。")
    replace_start(doc, "划分seed为", "主要分析的划分种子为11、23、37、53、71、83、97、113、127和149。分类采用种子化分层骨架组折，回归采用种子化骨架组平衡折；每个种子含3个外层骨架折，并在每个外层训练分区内生成3个内层骨架折。Bemis–Murcko骨架组保持完整，分配使用种子相关的随机并列处理并尽量平衡样本数，且不依据模型性能选择划分。仅使用前五个种子的分析均明确标注为五种子次要分析。")
    replace_start(doc, "所有单元层指标先在", "主要轻量候选库估计量先在每个划分种子内平均3个外层折，再在每个终点内以10个种子为区组汇总。五种子次要面板保留原分析单元并明确标注。跨终点描述性汇总使用终点层中位数、四分位距和范围；分类与回归不合并，因为ROC-AUC差与RMSE增加不是可交换单位。")
    replace_start(doc, "对每个终点、K和留出seed", "对每个端点和每个留出划分种子，使用其余9个种子上全部32个候选的平均外层效用选择K不变完整候选库参考，并在留出种子的3个外层折上评估。该参考身份在K = 4、8、16、32之间固定；K依赖参考只在相应合资格前缀内选择，同折最大值仅作为描述性有限集合机会界。负的K = 32减K = 4完成差距对比表示K = 32下相对于完整候选库参考的完成差距更小，正值表示K = 4下更小。")
    replace_start(doc, "令 u = (s, f)", "令u = (s, f)表示由划分种子s与外层折f构成的外层审计单元；V(u, j)为候选j的平均内层验证效用，A(u, j)为其外层审计效用。主要分析S_main = 10且F = 3；明确标注的次要分析S_secondary = 5。符合资格的前缀C_K、候选顺序与并列规则均在计算报告对比前固定。")
    replace_start(doc, "审计估计量使用的数学符号", "表3. 审计估计量使用的数学符号。")
    replace_start(doc, "公式（1）", "公式（1）–（4）定义候选选择、有限审计选择损失和机会校正排序；公式（5）–（6）定义交叉拟合参考与差距；公式（7）–（10）定义矩阵变换和有效秩；公式（11）–（13）定义候选库组成对比与归一化增益；公式（14）定义归一化选择熵。")
    replace_start(doc, "表示除留出种子", "参考训练集包含除留出种子s以外的全部外层单元。因此，主要K不变与K依赖参考使用S_main = 10中的其余9个种子；五种子次要分析使用S_secondary = 5中的其余4个种子。公式（5）–（6）定义交叉拟合身份与差距，可用性分解见第2.8节。")
    replace_start(doc, "X 为外层效用矩阵", "X为外层效用矩阵；Ledoit–Wolf表达式中的S_cov表示样本协方差，不得与S_main = 10或S_secondary = 5混淆。T为缩放单位阵目标，λ_i为收缩相关矩阵的非负特征值，p_i为其和为1的比例。公式（7）–（10）定义矩阵变换与有效秩。")
    replace_start(doc, "数值稳定常数", "数值稳定常数取ε = 10⁻¹²。Ledoit–Wolf收缩采用缩放单位阵目标及scikit-learn的解析收缩系数估计；特征分解前将收缩协方差重新缩放为相关矩阵R，仅因浮点误差产生的微小负特征值裁剪为0，并约定0 log 0 := 0。配对归一化增益与交叉拟合差距先在配对外层审计单元内计算，随后在每个种子内平均3个外层折，再以种子为区组进行端点内汇总；绝对分母不超过ε的单元记为缺失并单独报告。")
    replace_start(doc, "对每个终点和K构建候选效用矩阵",
        "十种子主要有效多样性审计对每个端点和K构建30 × K外层效用矩阵和90 × K内层效用矩阵，来源为10个划分种子、每种子3个外层折以及每个外层训练分区3个内层折。预设分析原始效用、逐行中心化效用、固定参照相对效用和单元内秩四种变换。图2与表S6展示十种子外层矩阵；表S6同时报告相应内层矩阵。五种子候选库组成与多表征矩阵仍作为独立次要分析，不与上述估计合并。")
    replace_start(doc, "报告经验相关谱",
        "有效多样性由候选效用矩阵而非候选标签或模型家族数量计算。报告经验相关谱、Ledoit–Wolf收缩相关谱、谱熵有效秩、参与率有效秩、非对角相关中位数及收缩系数。数值方差不超过10⁻¹²的候选列在相关估计前剔除并记录变换后列数；固定参照相对矩阵以候选1为参照，因此含K − 1列。敏感性分析逐一删除10个种子、外层折标签以及内层矩阵的内层折标签，并用3个额外的候选库预设参照重复计算；参照不依据外层表现选择。跨端点汇总采用中位数、IQR和范围，不把9个端点视为总体随机样本，也不报告总体推断性置信区间（表S6–S7）。")
    replace_start(doc, "在5个分类端点上，将ROC-AUC选择",
        "在5个分类端点的冻结十种子嵌套划分上记录内层ROC-AUC、PR-AUC和操作点指标。每个内层折以其训练分区中频数较少的类别定义少数类，并枚举内层验证预测产生的全部不同ROC阈值。若至少一个阈值达到少数类召回≥0.80，则先最大化多数类召回；并列时依次选择更高的少数类召回、更高的平衡召回以及更接近0.5的阈值。若没有阈值达到0.80，则在所有少数类召回有限的阈值上应用同一并列规则，并记录阈值不可行。候选的最终阈值取3个内层折阈值的中位数，少数类标签取众数，不实施事后概率校准。候选选择先保留平均内层少数类召回≥0.80的候选；若无候选合格，则保留完整合资格前缀。随后最大化平均内层多数类召回，并依次以平均内层少数类召回、平均内层PR-AUC和预注册顺序解决并列。候选身份、聚合阈值和少数类标签在一次性外层折评价前冻结；报告阈值可行性、候选资格、外层目标满足、少数类与多数类召回、少数类漏检率、PR-AUC和ROC-AUC。")
    replace_start(doc, "K = 32时，九个端点的Ledoit–Wolf谱熵秩中位数",
        f"十种子主要外层有效多样性审计中，K = 32时九个端点的Ledoit–Wolf谱熵秩中位数分别为：原始效用{div['raw']:.2f}、逐行中心化{div['row_centred']:.2f}、固定参照相对效用{div['fixed_reference_relative']:.2f}、单元内秩{div['within_unit_rank']:.2f}。四种矩阵分别保留共同审计难度、去除共同水平位移、表达相对于预设参照的差异或仅保留顺序。30 × 32外层矩阵和90 × 32内层矩阵均不与五种子次要候选库组成分析合并（图2；表S6–S7）。")
    replace_start(doc, "估计器、种子、折和参照敏感性",
        "矩阵层级、估计器、留一种子、留一折及参照敏感性会改变数值幅度，但名义K与矩阵依赖有效秩仍是不同量。完整IQR、范围、参与率秩、收缩系数、相关性摘要及内层矩阵结果见表S6；留一单元与预设参照结果见表S7。")
    replace_start(doc, "图2.",
        "图2. 十种子矩阵依赖的候选多样性。A，比较原始、逐行中心化、固定参照相对和单元内秩外层效用矩阵的Ledoit–Wolf谱熵秩；线为端点中位数，阴影为端点IQR，灰色点线为名义K。B，K = 32时谱熵秩与参与率秩的一致性及45°参考线。C，校正后的候选相关中位数。D，K = 32固定参照相对端点估计；实线和虚线范围分别删除一个种子与一个外层折标签，青绿色实心点为完整十种子估计，空心菱形使用预设线性参照。每个外层矩阵含30个种子–折行；相应90行内层矩阵结果见表S6–S7。")
    replace_start(doc, "图6.",
        "图6. 化学支持边界处的预测可靠性。A，在对称四模型矩阵中同时展示预测相关与高误差Jaccard重叠，并使用图内独立色带。B，将分类ROC-AUC、回归RMSE、分类假阴性率以及分类/回归高误差富集整合为Tanimoto支持风险矩阵；单元格文字为自然尺度中位数，颜色仅编码每行的不利方向。C，在对数比值轴上展示新骨架相对于已见或相关骨架的错误重叠、模型分歧、高误差富集和假阴性富集变化。D，分开展示四个固定配置模型的区分度与少数类安全性指标；紫色虚线为0.90覆盖率目标，预测集大小仅为显示而由1–2映射至0–1。")
    replace_start(doc, "图7.",
        "图7. 扩展候选池组成干预。A，K = 32时采用配对同质候选池审计最佳归一化，并展示六个预设端点的选择值与有限审计最佳机会。B，分别在分类和回归任务区段展示K = 4、8、16和32下的验证选择增益与交叉拟合差距，候选池图例仅显示一次。C，展示各端点、候选池和K组合的CAHit@3，负值完整保留；主图最右侧Entropy列报告K = 32的标准化候选选择熵。H表示同质Morgan，MV表示经典多表征，M表示现代增强。D，比较等K圆点与等下游预算菱形条件下的标准化选择增益和下游拟合/预测时间；标记面积表示K，并展示经验Pareto前沿。该组成干预为五种子次要分析，先在种子内平均3个外层折；时间不含模型获取、编码器预训练和缓存嵌入提取。")
    replace_start(doc, "交叉拟合K = 32减K = 4效应", "相对于K不变的完整候选库参考，9个端点中有8个端点的K = 32减K = 4完成差距对比为负，1个端点为正；其中7个端点的描述性划分种子敏感性区间不含0。负值表示K = 32下相对于完整候选库参考的完成差距更小，正值表示K = 4下更小。图3C、表4与本节由同一机器可读源表生成，不合并分类ROC-AUC与回归RMSE尺度。")
    replace_start(doc, "表3. 候选池扩张的交叉拟合效应", "表4. K不变完整候选库完成差距对比。")
    replace_start(doc, "对每个端点和每个留出划分种子，先用其余9个种子", "对每个端点和每个留出划分种子，先用其余9个种子上全部32个候选的平均外层效用选择一次参考候选，并以预注册顺序解决并列。该候选身份在留出种子的各折上评估，并在K = 4、8、16和32之间保持不变，即使它在较小K下不具备被选择资格。负的K = 32减K = 4完成差距对比表示K = 32下相对于完整候选库参考的完成差距更小；正值表示K = 4下更小。该估计量隔离了参考选择与留出种子，但仍复用同一公开端点与划分生成器，不构成独立队列验证。")

    ai = next((p for p in doc.paragraphs if p.text.strip().startswith(("2.17 Use of generative", "2.17 生成式"))), None)
    methods_target = ai or next((p for p in doc.paragraphs if p.text.strip() in {"3 Results", "3 结果"}), None)
    if methods_target:
        if ai: ai.text = "2.21 生成式人工智能的使用"
        sections = [
            ("2.17 K不变的K = 32交叉拟合参考", "对每个端点和每个留出划分种子，先用其余9个种子上全部32个候选的平均外层效用选择一次参考候选，并以预注册顺序解决并列。该候选身份在留出种子的各折上评估，并在K = 4、8、16和32之间保持不变，即使它在较小K下不具备被选择资格。负的K = 32减K = 4完成差距对比表示K = 32下相对于完整候选库参考的完成差距更小；正值表示K = 4下更小。该估计量隔离了参考选择与留出种子，但仍复用同一公开端点与划分生成器，不构成独立队列验证。"),
            ("2.18 实用等价与近等价候选集", "分类任务回顾性锁定的ROC-AUC容差网格为0.005、0.010和0.020，回归任务RMSE容差网格为0.025、0.050和0.100。若所选候选与交叉拟合参考的差距不超过ε，则记为ε成功；同时报告交叉拟合近等价候选集大小。所有阈值分析均作为敏感性描述，不作为事后显著性检验。"),
            ("2.19 PR-AUC与少数类召回约束选择", "在5个分类端点的冻结十种子嵌套划分上记录内层ROC-AUC、PR-AUC和操作点指标。每个内层折以其训练分区中频数较少的类别定义少数类，并枚举内层验证预测产生的全部不同ROC阈值。若至少一个阈值达到少数类召回≥0.80，则先最大化多数类召回；并列时依次选择更高的少数类召回、更高的平衡召回以及更接近0.5的阈值。若没有阈值达到0.80，则在所有少数类召回有限的阈值上应用同一并列规则，并记录阈值不可行。候选最终阈值取3个内层折阈值的中位数，少数类标签取众数，不实施事后概率校准。候选选择先保留平均内层少数类召回≥0.80的候选；若无候选合格，则保留完整合资格前缀。随后最大化平均内层多数类召回，并依次以平均内层少数类召回、平均内层PR-AUC和预注册顺序解决并列。候选身份、聚合阈值和少数类标签在一次性外层折评价前冻结；报告阈值可行性、候选资格、外层目标满足、少数类与多数类召回、少数类漏检率、PR-AUC和ROC-AUC。"),
            ("2.20 构造对照、顺序敏感性与恢复模拟", "以重复候选、弱候选和互补候选构造依赖结构，并评估预注册前缀、随机顺序、随机子集及家族平衡组成。经验相关异方差模拟复用观测到的候选相关结构和噪声尺度，检验不同信号强度与候选数量下的排序恢复。五种子组成与可靠性结果始终标记为次要分析。"),
        ]
        if not ai:
            sections.append(("2.21 生成式人工智能的使用", "生成式人工智能工具仅用于语言润色、结构检查和代码辅助；所有分析定义、数值、图表、引文与科学判断均由作者核验。工具未被用来生成或替代实验数据。"))
        for h, b in sections: insert_before(methods_target, h, b, doc)

    replace_start(doc, "端点特异的K = 32减K = 4", "端点特异的K = 32减K = 4对比先在10个划分种子内平均3个外层折，再对种子区组进行10 000次重采样。分类采用种子化分层骨架组折，回归在全部10次当前运行中采用种子化骨架组平衡折。历史无种子的回归GroupKFold仅用于划分转换审计，不混入主要估计量。重采样降低蒙特卡洛误差，但不会增加10个种子所提供的信息量；95%区间为描述性划分种子敏感性区间，不是针对未见分子端点的总体置信区间。分类ROC-AUC与回归RMSE不合并，且未定义跨端点汇总P值或确认性多重性族。")
    replace_start(doc, "图1.", "图1. 嵌套审计与证据层级。三个内层折用于排序预注册候选，三个外层折用于审计冻结决策。轻量候选库使用10个种子化骨架划分；组成与可靠性等次要面板保留其原预设分析单元。K不变K = 32参考、K依赖交叉拟合参考与描述性同折机会上界明确分离，并结合实用等价集合、指标匹配选择、构造候选对照、恢复模拟、化学支持与有界下游计算。")
    replace_start(doc, "图3 ", "图3. 机会校正排序校准与K不变完整候选库效应。A，五种子次要分析的端点中位CAHit@3、标准化MRR增益、端点IQR及95%随机排序包络。B，五种子次要正对照在6级注入验证—审计信号与4种候选数量下的恢复。C，十种子主要分析相对K不变K = 32交叉拟合参考的端点特异K = 32减K = 4完整候选库完成差距对比；分类ROC-AUC与回归RMSE使用独立坐标轴，负值表示K = 32差距更小，实心标记表示敏感性区间不含0；数值与表4完全一致。D，五种子次要组成对照；K = 32时各模式按设计等于完整候选库。")

    discussion = next((p for p in doc.paragraphs if p.text.strip() in {"4 Discussion", "4 讨论"}), None)
    if discussion:
        sections = [
            ("3.11 实用等价与指标依赖选择", "固定K = 32交叉拟合参考身份可避免随K改变比较对象。K = 32时，分类在ROC-AUC容差0.005、0.010和0.020下的ε成功率分别为62.7%、80.0%和86.0%；回归在RMSE容差0.025、0.050和0.100下分别为73.3%、79.2%和82.5%。任务匹配中间容差的合并描述性成功率为79.6%，平均交叉拟合近等价集合大小为10.17，所选候选属于该集合的比例为85.9%，验证集与交叉拟合集合的平均Jaccard重叠为0.490。这些均为回顾性锁定的报告容差，不是临床或药理学阈值。K = 32时，PR-AUC规则在52.7%的单元中改变ROC-AUC所选候选，并使外层PR-AUC、ROC-AUC、少数类召回、多数类召回和少数类漏检率平均变化+0.0018、−0.0026、−0.0002、−0.0136和+0.0002。少数类召回约束规则在66.0%的单元中改变候选，并使上述指标平均变化+0.0007、−0.0008、−0.0094、+0.0059和+0.0094。K = 32时100.0%的单元存在内层合资格候选，但冻结规则仅在62.0%的外层单元达到少数类召回≥0.80；负的平均召回变化被完整保留，不解释为约束可自动迁移（表S38–S40）。"),
            ("3.12 构造依赖对照与经验恢复", "重复和弱候选扩大名义K而未相应增加有效机会，互补候选的作用则依赖端点和信号强度。相关异方差模拟表明，有限验证下的恢复率同时受候选相关性、噪声异质性和候选数量影响；负结果完整保留在补充材料中。"),
            ("3.13 次要背景与负结果", "五种子候选库组成干预、计算暴露、化学支持与结构分离结果用于界定主要结论的适用范围，不与十种子主要估计量合并。所有未观察到稳定改善或未达到约束目标的结果均在正文概述，并在补充文件中提供机器可读细节。"),
        ]
        for h, b in sections: insert_before(discussion, h, b, doc)
    limitations = next((p for p in doc.paragraphs if p.text.strip().startswith(("4.9 Limitations", "4.9 局限"))), None)
    if limitations:
        limitations.text = "4.11 局限性"
        insert_before(limitations, "4.9 实用等价与指标匹配选择", "实用等价集合可将注意力从不稳定的单一冠军转向在预先声明容差内可互换的候选，但容差具有任务尺度，必须在看到结果前锁定。PR-AUC或少数类召回规则只改变选择目标，并不保证外层泛化约束被满足。", doc)
        insert_before(limitations, "4.10 候选依赖与有限验证", "构造对照和恢复模拟共同说明，新增候选的数量、独立信息量与验证可分辨性是三个不同概念。报告名义K时应同时描述候选相关结构、顺序、可用性以及验证不确定性。", doc)

    # Keep Chinese captions explicit while retaining the same figure content and evidence.
    for n in range(1, 9):
        cap = find_caption(doc, f"Figure {n}.")
        if cap: cap.text = cap.text.replace(f"Figure {n}.", f"图{n}.", 1)
        cap = find_figure_caption(doc, n)
        if cap and not cap.text.strip().startswith(f"图{n}."):
            cap.text = re.sub(rf"^图\s*{n}(?:[.．]|\s)*", f"图{n}. ", cap.text.strip(), count=1)
        replace_figure(doc, n)
    replace_start(doc, "图2.",
        "图2. 十种子矩阵依赖的候选多样性。A，比较原始、逐行中心化、固定参照相对和单元内秩外层效用矩阵的Ledoit–Wolf谱熵秩；线为端点中位数，阴影为端点IQR，灰色点线为名义K。B，K = 32时谱熵秩与参与率秩的一致性及45°参考线。C，校正后的候选相关中位数。D，K = 32固定参照相对端点估计；实线和虚线范围分别删除一个种子与一个外层折标签，青绿色实心点为完整十种子估计，空心菱形使用预设线性参照。每个外层矩阵含30个种子–折行；相应90行内层矩阵结果见表S6–S7。")
    if find_caption(doc, "图8.") is None:
        target = discussion or next((p for p in doc.paragraphs if p.text.strip().startswith("4 ")), None)
        if target:
            pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER; pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(str(FIG / "Figure8_600dpi.png"), width=Inches(6.45))
            cap = doc.add_paragraph("图8. 十种子主要审计中的实用等价与指标依赖选择。A面板上下排列分类ROC-AUC和回归RMSE差距，使用独立纵轴并共享候选数K横轴；圆点/实线、方块/虚线和三角/点线分别表示K不变、K依赖和同折估计量，使灰度打印时仍可区分。B，展示回顾性锁定中间容差下各K的交叉拟合ε成功率。C，在单一纵轴上展示交叉拟合近等价集合的平均大小，B中的分类/回归图例同时适用于C。D，明确标示五个分类端点的PR-AUC候选切换、召回规则候选切换、ΔPR-AUC和Δ少数类召回；少数类召回负结果完整保留。A–D按2 × 2四联图排列。", style="Figure Caption")
            target._p.addprevious(pic._p); target._p.addprevious(cap._p)
    replace_start(doc, "Additional file 1：", "Additional file 1：补充方法、结果与补充表说明。Additional file 2：机器可读补充表S1–S46。Additional file 3：补充图S1–S25。Additional file 4：代码、环境、清单与复现说明。")
    replace_start(doc, "本研究使用的公开数据来源", "本研究所用公开数据源见Additional file 2的Table S1。最新核验的公开代码位于https://github.com/zfr0857/FZYC-Mol，release为paper-release-2026-07-r9，固定commit为9635a902fa3cc7bb7b71a234c1b2bbbe415193f0。Additional file 4包含该可移植基础和发布后新增的paper43十种子完成层；后者尚未进入公开tag，因此仓库同步及不可变归档仍是投稿阻断项，本版本未声称归档DOI。")
    revise_tables(doc, chinese=True); apply_native_inline_math(doc); add_supplement_statement(doc, chinese=True); style_paragraphs(doc, chinese=True); keep_figures_with_captions(doc)
    doc.core_properties.title = doc.paragraphs[0].text
    doc.save(OUT_ZH)


def main() -> None:
    format_english(); build_chinese(); print(OUT_EN); print(OUT_ZH)


if __name__ == "__main__": main()
