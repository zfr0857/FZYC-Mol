from __future__ import annotations

import csv
import os
import re
import shutil
import tempfile
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from openpyxl import load_workbook


ROOT = Path(r"D:\fzyc")
PKG = ROOT / "output" / "paper43_jcheminform_completion_20260726"
AUTHOR = PKG / "author_review"
QC = PKG / "quality_control_final"
UPLOAD = PKG / "upload_ready"

EN = AUTHOR / "Main_manuscript_Journal_of_Cheminformatics_FINAL_CLEAN.docx"
EN_OLD = AUTHOR / "Main_manuscript_Journal_of_Cheminformatics_PRE_FINAL_ROUND_BACKUP.docx"
EN_REVISED = AUTHOR / "Main_manuscript_Journal_of_Cheminformatics_AUTHOR_REVIEW_CONFLICT_HOLD.docx"
ZH = AUTHOR / "候选池扩张_有限验证下的机会与选择稳定性_中文同步终稿.docx"
ZH_OLD = AUTHOR / "候选池扩张_有限验证下的机会与选择稳定性_修订前备份.docx"
ZH_REVISED = AUTHOR / "候选池扩张_有限验证下的机会与选择稳定性_中文同步终稿_CONFLICT_HOLD.docx"
SUPP = AUTHOR / "Additional_file_1_Supplementary_Methods_and_Results_FINAL.docx"
SUPP_OLD = AUTHOR / "Additional_file_1_Supplementary_Methods_and_Results_PRE_FINAL_ROUND_BACKUP.docx"
XLSX = UPLOAD / "Additional_file_2_Machine_readable_Supplementary_Tables_S1-S46.xlsx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = {"w": W, "m": M}


def first(doc: Document, prefix: str):
    return next((p for p in doc.paragraphs if p.text.strip().startswith(prefix)), None)


def set_paragraph(paragraph, text: str) -> None:
    paragraph.text = text


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_before(paragraph, text: str, style: str | None = None):
    node = OxmlElement("w:p")
    paragraph._p.addprevious(node)
    from docx.text.paragraph import Paragraph
    inserted = Paragraph(node, paragraph._parent)
    if style:
        inserted.style = style
    inserted.add_run(text)
    return inserted


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        try:
            paragraph.style = "Table Text"
        except KeyError:
            pass
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mrun(text: str):
    run = etree.Element(etree.QName(M, "r"))
    rpr = etree.SubElement(run, etree.QName(M, "rPr"))
    style = etree.SubElement(rpr, etree.QName(M, "sty"))
    style.set(etree.QName(M, "val"), "p")
    node = etree.SubElement(run, etree.QName(M, "t"))
    node.text = text
    return run


def as_nodes(value):
    if value is None:
        return []
    if isinstance(value, str):
        return [mrun(value)]
    if isinstance(value, (list, tuple)):
        result = []
        for item in value:
            result.extend(as_nodes(item))
        return result
    return [value]


def mslot(tag: str, value):
    slot = etree.Element(etree.QName(M, tag))
    for node in as_nodes(value):
        slot.append(deepcopy(node))
    return slot


def msub(base, subscript):
    node = etree.Element(etree.QName(M, "sSub"))
    node.append(mslot("e", base))
    node.append(mslot("sub", subscript))
    return node


def msup(base, superscript):
    node = etree.Element(etree.QName(M, "sSup"))
    node.append(mslot("e", base))
    node.append(mslot("sup", superscript))
    return node


def msubsup(base, subscript, superscript):
    node = etree.Element(etree.QName(M, "sSubSup"))
    node.append(mslot("e", base))
    node.append(mslot("sub", subscript))
    node.append(mslot("sup", superscript))
    return node


def macc(base):
    node = etree.Element(etree.QName(M, "acc"))
    pr = etree.SubElement(node, etree.QName(M, "accPr"))
    char = etree.SubElement(pr, etree.QName(M, "chr"))
    char.set(etree.QName(M, "val"), "̂")
    node.append(mslot("e", base))
    return node


def mfrac(numerator, denominator):
    node = etree.Element(etree.QName(M, "f"))
    node.append(mslot("num", numerator))
    node.append(mslot("den", denominator))
    return node


def msum(subscript, superscript=None):
    node = etree.Element(etree.QName(M, "nary"))
    pr = etree.SubElement(node, etree.QName(M, "naryPr"))
    char = etree.SubElement(pr, etree.QName(M, "chr"))
    char.set(etree.QName(M, "val"), "∑")
    lim = etree.SubElement(pr, etree.QName(M, "limLoc"))
    lim.set(etree.QName(M, "val"), "undOvr")
    node.append(mslot("sub", subscript))
    node.append(mslot("sup", superscript or ""))
    node.append(mslot("e", ""))
    return node


def math_object(items):
    obj = etree.Element(etree.QName(M, "oMath"))
    for node in as_nodes(items):
        obj.append(deepcopy(node))
    return obj


def eq_array(rows):
    array = etree.Element(etree.QName(M, "eqArr"))
    for row in rows:
        array.append(mslot("e", row))
    return math_object(array)


def jhat():
    return [msub(macc("j"), "u"), mrun("(K)")]


def jref(kind: str, args: str):
    return [msubsup("j", "ref", kind), mrun(f"({args})")]


def gsymbol(kind: str):
    return msub("G", kind)


def csymbol(k: str):
    return msub("C", k)


def usymbol(sub: str):
    return msub("U", sub)


def selected_formula():
    return math_object([jhat(), " ≡ ", msubsup("j", "u", "K"), " = ", msub("arg max", ["j∈", csymbol("K")]), " V(u,j)"])


def reference_formula(kind: str):
    k = "32" if kind == "inv" else "K"
    return math_object([
        jref(kind, "−s" if kind == "inv" else "−s,K"), " = ",
        msub("arg max", ["j∈", csymbol(k)]), " ",
        mfrac("1", ["|", usymbol("−s"), "|"]), " ",
        msum(["u∈", usymbol("−s")]), " A(u,j)",
    ])


def gaps_formula():
    return eq_array([
        [gsymbol("inv"), "(u,K) = A(u,", jref("inv", "−s"), ") − A(u,", jhat(), ")"],
        [gsymbol("dep"), "(u,K) = A(u,", jref("dep", "−s,K"), ") − A(u,", jhat(), ")"],
    ])


def same_formula():
    return math_object([
        gsymbol("same"), "(u,K) = ", msub("max", ["j∈", csymbol("K")]), " A(u,j) − A(u,", jhat(), ")"
    ])


def availability_formula():
    return eq_array([
        [gsymbol("avail"), "(u,K) = A(u,", jref("inv", "−s"), ") − A(u,", jref("dep", "−s,K"), ")"],
        [gsymbol("inv"), "(u,K) = ", gsymbol("avail"), "(u,K) + ", gsymbol("dep"), "(u,K)"],
    ])


def delta_formula():
    return math_object([
        msub("Δ", "inv"), "(e) = ", mfrac("1", [msub("S", "main"), "F"]), " ",
        msum(["s∈", msup("S", "e")]), " ", msum("f=1", "F"), " [",
        gsymbol("inv"), "((s,f),32) − ", gsymbol("inv"), "((s,f),4)]",
    ])


def equation_paragraph(template, math, number: int):
    paragraph = deepcopy(template)
    for child in list(paragraph):
        if child.tag != etree.QName(W, "pPr"):
            paragraph.remove(child)
    ppr = paragraph.find("w:pPr", NS)
    if ppr is None:
        ppr = etree.Element(etree.QName(W, "pPr"))
        paragraph.insert(0, ppr)
    tabs = ppr.find("w:tabs", NS)
    if tabs is None:
        tabs = etree.SubElement(ppr, etree.QName(W, "tabs"))
        for val, pos in [("center", "4649"), ("right", "9298")]:
            tab = etree.SubElement(tabs, etree.QName(W, "tab"))
            tab.set(etree.QName(W, "val"), val)
            tab.set(etree.QName(W, "pos"), pos)
    left_tab = etree.SubElement(paragraph, etree.QName(W, "r"))
    etree.SubElement(left_tab, etree.QName(W, "tab"))
    paragraph.append(math)
    right_tab = etree.SubElement(paragraph, etree.QName(W, "r"))
    etree.SubElement(right_tab, etree.QName(W, "tab"))
    number_run = etree.SubElement(paragraph, etree.QName(W, "r"))
    text = etree.SubElement(number_run, etree.QName(W, "t"))
    text.text = f"({number})"
    return paragraph


def replace_zip_members(path: Path, replacements: dict[str, bytes]) -> None:
    handle, name = tempfile.mkstemp(suffix=path.suffix, dir=path.parent)
    os.close(handle)
    temp = Path(name)
    try:
        with ZipFile(path) as source, ZipFile(temp, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                target.writestr(item, replacements.get(item.filename, source.read(item.filename)))
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def patch_equations(path: Path) -> None:
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    body = root.find("w:body", NS)
    records = {}
    for paragraph in body.xpath("./w:p[.//m:oMath]", namespaces=NS):
        text = "".join(paragraph.xpath(".//w:t/text() | .//m:t/text()", namespaces=NS))
        match = re.search(r"\((\d+)\)\s*$", text)
        if match:
            records[int(match.group(1))] = paragraph
    if set(records) != set(range(1, 15)):
        raise RuntimeError(f"Expected equations 1–14 in {path.name}; found {sorted(records)}")
    template = records[5]
    index = body.index(records[5])
    body.remove(records[5])
    body.remove(records[6])
    new_math = [
        selected_formula(), reference_formula("inv"), reference_formula("dep"), gaps_formula(),
        same_formula(), availability_formula(), delta_formula(),
    ]
    # The seven required estimand blocks occupy Equations (5)–(11).
    for offset, math in enumerate(new_math):
        body.insert(index + offset, equation_paragraph(template, math, 5 + offset))
    shift = 5
    for old in range(7, 15):
        paragraph = records[old]
        number_nodes = paragraph.xpath("./w:r[last()]/w:t", namespaces=NS)
        if len(number_nodes) != 1:
            raise RuntimeError(f"Cannot renumber Equation {old} in {path.name}")
        number_nodes[0].text = f"({old + shift})"
    replace_zip_members(path, {"word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")})


def symbol_math(key: str):
    mapping = {
        "S_main": msub("S", "main"), "S_secondary": msub("S", "secondary"),
        "s": mrun("s"), "F": mrun("F"), "f": mrun("f"),
        "u=(s,f)": [mrun("u = (s,f)")], "j": mrun("j"), "K": mrun("K"),
        "C_K": msub("C", "K"), "V(u,j)": mrun("V(u,j)"), "A(u,j)": mrun("A(u,j)"),
        "jhat": jhat(), "jref_inv": jref("inv", "−s"), "jref_dep": jref("dep", "−s,K"),
        "G_inv": [gsymbol("inv"), mrun("(u,K)")], "G_dep": [gsymbol("dep"), mrun("(u,K)")],
        "G_same": [gsymbol("same"), mrun("(u,K)")], "G_avail": [gsymbol("avail"), mrun("(u,K)")],
        "Delta_inv": [msub("Δ", "inv"), mrun("(e)")],
    }
    return math_object(mapping[key])


def set_symbol_cell(cell, key: str) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    try:
        paragraph.style = "Table Text"
    except KeyError:
        pass
    paragraph._p.append(symbol_math(key))
    paragraph.paragraph_format.keep_together = True
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def resize_table2(table) -> None:
    widths = [Inches(1.75), Inches(1.05), Inches(2.05), Inches(1.60)]
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.w = width
    for row_index, row in enumerate(table.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = width
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.type = "dxa"
            tcw.w = width.twips
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)


def rebuild_table3(doc: Document, chinese: bool) -> None:
    rows = [
        ("header", "符号" if chinese else "Symbol", "定义" if chinese else "Definition"),
        ("S_main", "", "主要分析的划分种子数（10）" if chinese else "number of primary split seeds (10)"),
        ("S_secondary", "", "明确标注的次要分析所用种子数（5）" if chinese else "number of seeds in explicitly labelled secondary analyses (5)"),
        ("s", "", "划分种子" if chinese else "split seed"),
        ("F", "", "每个种子的外层折数（3）" if chinese else "number of outer folds per seed (3)"),
        ("f", "", "外层折" if chinese else "outer fold"),
        ("u=(s,f)", "", "外层审计单元" if chinese else "outer audit unit"),
        ("j", "", "候选" if chinese else "candidate"),
        ("K", "", "候选数量" if chinese else "candidate count"),
        ("C_K", "", "规模K下符合资格的预注册候选前缀" if chinese else "eligible registered candidate prefix at size K"),
        ("V(u,j)", "", "内层验证效用" if chinese else "inner-validation utility"),
        ("A(u,j)", "", "外层审计效用（回归使用负RMSE方向）" if chinese else "outer-audit utility (negative RMSE direction for regression)"),
        ("jhat", "", "规模K下由内层验证选择的候选" if chinese else "validation-selected candidate at size K"),
        ("jref_inv", "", "不使用种子s选出的K不变完整候选库参考" if chinese else "K-invariant full-registry reference selected without seed s"),
        ("jref_dep", "", "不使用种子s、在合资格前缀内选出的K依赖参考" if chinese else "K-dependent eligible-prefix reference selected without seed s"),
        ("G_inv", "", "K不变完整候选库完成差距" if chinese else "K-invariant full-registry completion gap"),
        ("G_dep", "", "K依赖合资格前缀选择差距" if chinese else "K-dependent within-prefix selection gap"),
        ("G_same", "", "同折有限集合机会差距" if chinese else "same-fold finite-set opportunity gap"),
        ("G_avail", "", "可用性分量；G_inv = G_avail + G_dep" if chinese else "availability component; G_inv = G_avail + G_dep"),
        ("Delta_inv", "", "端点层K=32减K=4的K不变完成差距对比" if chinese else "endpoint-level K=32 minus K=4 K-invariant completion-gap contrast"),
    ]
    table = doc.tables[2]
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for i, (key, label, definition) in enumerate(rows):
        if i == 0:
            set_cell_text(table.cell(i, 0), label, True)
            set_cell_text(table.cell(i, 1), definition, True)
        else:
            set_symbol_cell(table.cell(i, 0), key)
            set_cell_text(table.cell(i, 1), definition)
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Inches(1.55)
        row.cells[1].width = Inches(4.90)


def update_common_equation_text(doc: Document, chinese: bool) -> None:
    if chinese:
        replacements = {
            "公式（1）–（4）定义": "公式（1）–（4）定义候选选择、有限审计选择损失和机会校正排序；公式（5）–（11）定义验证所选候选、K不变与K依赖参考、四类差距、可用性分解及端点层Δ_inv对比；公式（12）–（15）定义矩阵变换和有效秩；公式（16）–（19）定义候选库组成对比、归一化增益和选择熵。",
            "参考训练集包含除留出种子": "参考训练集包含除留出种子s以外的全部外层单元。主要参考使用S_main = 10中的其余9个种子；五种子次要分析只在其自身范围内使用其余4个种子。公式（5）–（11）完整定义参考候选、G_inv、G_dep、G_same、G_avail、分解关系和Δ_inv(e)。负Δ_inv(e)表示K = 32下相对于完整候选库参考的完成差距更小；正值表示K = 4下更小。",
            "X为外层效用矩阵": "X为外层效用矩阵；Ledoit–Wolf表达式中的S_cov表示样本协方差，不得与S_main = 10或S_secondary = 5混淆。T为缩放单位阵目标，λ_i为收缩相关矩阵的非负特征值，p_i为其和为1的比例。公式（12）–（15）定义矩阵变换与有效秩。",
            "hom表示同质Morgan候选池": "hom表示同质Morgan候选池；完整下游效率与Pareto定义见补充方法。公式（16）–（19）定义组成对比、归一化增益和标准化选择熵。",
            "此处 π_j": "此处π_j为候选j在重复外层审计单元中被选中的比例，与CAHit@q定义中的排序截断q不同。公式（19）定义标准化选择熵。",
        }
    else:
        replacements = {
            "Equations (1)": "Equations (1)–(4) define candidate selection, finite-audit selection loss and chance-adjusted ranking. Equations (5)–(11) define the validation-selected candidate, K-invariant and K-dependent references, the four gaps, their availability decomposition and the endpoint-level Δ_inv contrast. Equations (12)–(15) define matrix transformations and effective ranks; Equations (16)–(19) define composition contrasts, normalized gains and selection entropy.",
            "The reference-training set contains": "The reference-training set contains all outer units from seeds other than held-out seed s. Primary references use the other nine of S_main = 10 seeds; five-seed secondary analyses use the other four only within their own scope. Equations (5)–(11) define the candidate identities, G_inv, G_dep, G_same, G_avail, the exact decomposition and Δ_inv(e). A negative Δ_inv(e) means a smaller full-registry completion gap at K = 32; a positive value means a smaller gap at K = 4.",
            "X is the outer-utility matrix": "X is the outer-utility matrix; S_cov in the Ledoit–Wolf expression denotes sample covariance and must not be confused with S_main = 10 or S_secondary = 5. T is the scaled-identity target, λ_i are the non-negative eigenvalues of the shrinkage correlation matrix and p_i their unit-sum proportions. Equations (12)–(15) define the transformations and effective ranks.",
            "The label hom denotes": "The label hom denotes the homogeneous Morgan pool; complete downstream-efficiency and Pareto definitions are provided in Supplementary Methods. Equations (16)–(19) define the composition contrasts, normalized gains and normalized selection entropy.",
            "Here π_j is": "Here π_j is the proportion of repeated outer audit units in which candidate j is selected; it is distinct from the rank cutoff q in the CAHit@q definition. Equation (19) defines normalized selection entropy.",
        }
    for prefix, value in replacements.items():
        paragraph = first(doc, prefix)
        if paragraph:
            paragraph.text = value
    heading = first(doc, "交叉拟合参照" if chinese else "Cross-fitted reference")
    if heading:
        heading.text = "完成差距估计量与分解" if chinese else "Completion-gap estimands and decomposition"


def revise_english(source: Path, target: Path) -> None:
    doc = Document(source)
    first(doc, "2.8 ").text = "2.8 Leave-one-seed-out references and completion-gap decomposition"
    first(doc, "For each held-out seed s,").text = (
        "For outer unit u = (s,f), the validation-selected candidate ĵ_u(K), the K-invariant full-registry reference "
        "j_ref^inv(−s), the K-dependent within-prefix reference j_ref^dep(−s,K), the K-invariant completion gap G_inv(u,K), "
        "the K-dependent selection gap G_dep(u,K), the same-fold finite-set opportunity gap G_same(u,K) and the availability "
        "component G_avail(u,K) are defined formally in Equations (5)–(11). The exact identity G_inv(u,K) = G_avail(u,K) + "
        "G_dep(u,K) separates candidate availability from within-prefix selection. The endpoint contrast Δ_inv(e) averages paired "
        "G_inv differences over ten seeds and three outer folds; a negative K = 32 minus K = 4 contrast indicates a smaller "
        "full-registry completion gap at K = 32, whereas a positive contrast indicates a smaller gap at K = 4."
    )
    first(doc, "Reference selection was isolated").text = (
        "Reference identities were selected without the held-out seed and registry order resolved ties. The same-fold maximum is "
        "descriptive rather than deployable. Seed isolation reduces same-unit circularity but reuses the same public endpoints, "
        "candidate registry and split generator; the ten seed blocks are sensitivity replicates, not independent studies or external validation."
    )
    for prefix in ["2.17 K-invariant", "For each endpoint and held-out seed,"]:
        p = first(doc, prefix)
        if p:
            delete_paragraph(p)
    first(doc, "2.18 Practical").text = "2.17 Practical equivalence and near-equivalent sets"
    first(doc, "2.19 PR-AUC").text = "2.18 PR-AUC and minority-recall-constrained selection"
    first(doc, "2.20 Constructed").text = "2.19 Constructed controls, order sensitivity and recovery simulation"
    p = first(doc, "At the largest registry each endpoint supplied")
    if p:
        p.text = (
            "In the primary ten-seed analysis, each endpoint supplied 30 outer audit units. Because the largest candidate count "
            "K = 32 still slightly exceeded the number of audit units, the unshrunk empirical candidate-correlation matrix could "
            "remain rank deficient. Ledoit–Wolf shrinkage stabilized covariance directions but did not create new independent "
            "information. Spectral-entropy and participation-ratio ranks weight the eigenvalue spectrum differently, so both were "
            "reported with omission and prespecified-reference sensitivities."
        )
    lim = first(doc, "4.11 Limitations")
    if lim:
        cursor = lim
        following = []
        node = lim._p.getnext()
        while node is not None and len(following) < 5:
            from docx.text.paragraph import Paragraph
            if node.tag == qn("w:p"):
                paragraph = Paragraph(node, lim._parent)
                if paragraph.text.strip().startswith(("Supplementary package", "The supplementary package", "Declarations")):
                    break
                following.append(paragraph)
            node = node.getnext()
        for paragraph in following:
            delete_paragraph(paragraph)
        paragraphs = [
            "This study is a retrospective audit of public datasets, not a prospective study, independent external validation or deployment evaluation. The primary analysis contains ten split seeds and three outer folds per seed, but its uncertainty information mainly comes from ten seed blocks; resampling cannot increase the independent information supplied by those blocks.",
            "The K-invariant full-registry reference separates reference selection from the held-out seed, but it reuses the same public endpoints, candidate registry and split generator. It therefore reduces one source of circularity without replacing validation in an independent cohort or data source.",
            "The ε tolerances are retrospectively locked reporting tolerances rather than clinical, pharmacological or industrial minimum important differences. PR-AUC and training-defined minority-recall rules changed the inner objective, but the corresponding outer-fold constraints did not transfer reliably.",
            "Coverage of modern candidates was limited. Equal-budget analyses record downstream fitting and prediction time but exclude model acquisition, encoder pretraining, cached embedding extraction and complete end-to-end cost. Tanimoto-component transport covered only three endpoints and one similarity rule.",
            "Endpoints, seeds, folds, candidates and overlapping subsets are dependent. Constructed duplicate, near-duplicate, weak and complementary controls and empirical recovery simulations support a mechanism-based interpretation, but they do not establish a single causal effect of effective diversity on selection gaps."
        ]
        anchor = lim._p
        for text in paragraphs:
            node = OxmlElement("w:p")
            anchor.addnext(node)
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(node, lim._parent)
            paragraph.style = doc.styles["Normal"]
            paragraph.add_run(text)
            anchor = node
    title_fields = {
        "[AUTHOR": "[Author names and ORCID iDs not provided — CONFLICT HOLD]",
        "[Institutional": "[Institutional affiliations and full postal addresses not provided — CONFLICT HOLD]",
        "*Corresponding author:": "*Corresponding author: [name, email and postal address not provided — CONFLICT HOLD]",
        "Author confirmation required": "[Not provided — CONFLICT HOLD; no declaration inferred.]",
    }
    for prefix, value in title_fields.items():
        for paragraph in [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]:
            paragraph.text = value
    avail = first(doc, "The datasets supporting this article")
    if avail:
        avail.text = (
            "The datasets supporting this article are public and the processed audit tables are included in Additional files 1–4. "
            "The latest verified public code release is https://github.com/zfr0857/FZYC-Mol/releases/tag/paper-release-2026-07-r9 "
            "at commit 9635a902fa3cc7bb7b71a234c1b2bbbe415193f0. Additional file 4 contains that portable base plus the "
            "post-release paper43 ten-seed completion overlay. Because the overlay is not present in a verified public r10 tag, "
            "repository synchronization remains a submission blocker. No separate archival DOI was available at the time of this review."
        )
    table2 = doc.tables[1]
    set_cell_text(table2.cell(1, 3), "34 560 candidate fits; candidate-fitting time 8 130.16 s")
    set_cell_text(table2.cell(5, 3), "1 080 outer units; downstream fitting/prediction time 64 616.35 s")
    resize_table2(table2)
    rebuild_table3(doc, False)
    update_common_equation_text(doc, False)
    doc.save(target)
    patch_equations(target)


def revise_chinese(source: Path, target: Path) -> None:
    doc = Document(source)
    first(doc, "2.8 ").text = "2.8 留一种子参考与完成差距分解"
    p = first(doc, "对每个端点和每个留出划分种子")
    if p:
        p.text = (
            "令外层审计单元u = (s,f)。公式（5）–（11）正式定义验证所选候选ĵ_u(K)、K不变完整候选库参考j_ref^inv(−s)、"
            "K依赖合资格前缀参考j_ref^dep(−s,K)、K不变完成差距G_inv(u,K)、K依赖前缀内选择差距G_dep(u,K)、"
            "同折有限集合机会差距G_same(u,K)和可用性分量G_avail(u,K)。恒等式G_inv(u,K) = G_avail(u,K) + G_dep(u,K)"
            "将候选可用性与前缀内选择分开。端点层Δ_inv(e)对10个种子和每种子的3个外层折求配对平均；负值表示K = 32下相对于完整候选库参考的完成差距更小，正值表示K = 4下更小。"
        )
        after = p._p.getnext()
        if after is None or "参考身份" not in "".join(after.itertext()):
            node = OxmlElement("w:p")
            p._p.addnext(node)
            from docx.text.paragraph import Paragraph
            q = Paragraph(node, p._parent)
            q.style = doc.styles["Normal"]
            q.add_run("参考身份在不使用留出种子的数据上选择，并以预注册候选顺序解决并列；同折最大值仅为描述性机会界。该设计降低同单元循环性，但仍复用相同公开端点、候选登记表和划分生成器，不构成独立外部验证。十个种子区组是敏感性重复，而不是十项独立研究。")
    for prefix in ["2.17 K不变", "对每个端点和每个留出划分种子，先用"]:
        p = first(doc, prefix)
        if p:
            delete_paragraph(p)
    first(doc, "2.18 实用等价").text = "2.17 实用等价与近等价候选集"
    first(doc, "2.19 PR-AUC").text = "2.18 PR-AUC与少数类召回约束选择"
    first(doc, "2.20 构造对照").text = "2.19 构造对照、顺序敏感性与恢复模拟"
    for prefix in ["2.21 生成式人工智能", "生成式人工智能工具仅用于语言润色"]:
        p = first(doc, prefix)
        if p:
            delete_paragraph(p)
    p = first(doc, "最大登记表中每个终点只有15个外层行")
    if p:
        p.text = (
            "主要十种子分析中，每个终点包含30个外层审计单元。由于最大候选数K = 32仍略高于审计单元数，未经收缩的经验候选相关矩阵可能秩亏。"
            "Ledoit–Wolf收缩能够稳定协方差方向，但不会创造新的独立信息。谱熵秩与参与率秩对特征谱赋权不同，因此二者均与留一种子、留一折和预设参照敏感性结果共同报告。"
        )
    lim = first(doc, "4.11 局限性")
    if lim:
        following = []
        node = lim._p.getnext()
        while node is not None and len(following) < 5:
            from docx.text.paragraph import Paragraph
            if node.tag == qn("w:p"):
                paragraph = Paragraph(node, lim._parent)
                if paragraph.text.strip().startswith(("补充材料", "声明")):
                    break
                following.append(paragraph)
            node = node.getnext()
        for paragraph in following:
            delete_paragraph(paragraph)
        texts = [
            "本研究是公开数据上的回顾性审计，不是前瞻性研究、独立外部验证或部署评价。主要分析包含10个划分种子且每种子3个外层折，但不确定性信息主要来自10个种子区组；重采样不能增加这些区组所提供的独立信息。",
            "K不变完整候选库参考将参考选择与留出种子分离，但仍复用相同的公开端点、候选登记表和划分生成器。因此，该设计降低一种循环性，却不能替代独立队列或独立数据源验证。",
            "ε容差是回顾性锁定的报告容差，不是临床、药理或工业最小重要差异。PR-AUC和由训练数据定义的少数类召回规则改变了内层目标，但相应外层约束并未稳定迁移。",
            "现代候选的覆盖范围有限。等预算分析只记录下游拟合/预测时间，不含模型获取、编码器预训练、缓存嵌入提取和完整端到端成本。Tanimoto分量迁移仅覆盖3个端点和1种相似度规则。",
            "端点、种子、折、候选及重叠子集相互依赖。重复、近重复、弱质和互补候选构造对照及经验恢复模拟支持机制解释，但不能证明有效多样性对选择差距具有单一因果效应。",
        ]
        anchor = lim._p
        for text in texts:
            node = OxmlElement("w:p")
            anchor.addnext(node)
            from docx.text.paragraph import Paragraph
            q = Paragraph(node, lim._parent)
            q.style = doc.styles["Normal"]
            q.add_run(text)
            anchor = node
    fields = {
        "[作者姓名": "[作者姓名与ORCID未提供——CONFLICT HOLD]",
        "[机构": "[作者机构与完整邮寄地址未提供——CONFLICT HOLD]",
        "*通讯作者": "*通讯作者：[姓名、电子邮箱和邮寄地址未提供——CONFLICT HOLD]",
        "投稿前须": "[未提供——CONFLICT HOLD；未推断声明内容。]",
        "投稿前需": "[未提供——CONFLICT HOLD；未推断声明内容。]",
        "需作者确认": "[未提供——CONFLICT HOLD；未推断声明内容。]",
    }
    for prefix, value in fields.items():
        for paragraph in [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]:
            paragraph.text = value
    avail = first(doc, "本研究所用公开数据源")
    if avail:
        avail.text = (
            "本研究所用公开数据源见Additional file 2的Table S1。当前可核验的最新公开代码版本为"
            "https://github.com/zfr0857/FZYC-Mol/releases/tag/paper-release-2026-07-r9，固定commit为"
            "9635a902fa3cc7bb7b71a234c1b2bbbe415193f0。Additional file 4包含该公开基础及发布后新增的paper43十种子完成层。"
            "由于该完成层尚未出现在可核验的公开r10 tag中，仓库同步仍是投稿阻断项。本轮审阅时没有可用的独立归档DOI。"
        )
    table2 = doc.tables[1]
    set_cell_text(table2.cell(1, 3), "34 560次候选拟合；候选拟合时间8 130.16 s")
    set_cell_text(table2.cell(5, 3), "1 080个外层单元；下游拟合/预测时间64 616.35 s")
    resize_table2(table2)
    rebuild_table3(doc, True)
    update_common_equation_text(doc, True)
    doc.save(target)
    patch_equations(target)


def revise_supplement(path: Path) -> None:
    doc = Document(path)
    p = first(doc, "Raw, row-centred")
    if p:
        p.text = (
            "The ten-seed primary effective-diversity analysis used raw, row-centred, fixed-reference-relative and within-unit-rank matrices. "
            "At K = 32, endpoint-median Ledoit–Wolf entropy ranks were 2.36, 19.25, 5.56 and 24.33, respectively. These values reflect "
            "matrix construction rather than independent candidate counts. Tables S6–S7 provide complete 30 × K outer and 90 × K inner "
            "results, leave-one-seed/fold sensitivities and prespecified-reference analyses; no five-seed effective-rank row is primary."
        )
    p = first(doc, "Equation-to-code mapping")
    if p:
        p.text = (
            "Equation-to-code mapping and version provenance. Main-text Equations (5)–(11) define ĵ_u(K), j_ref^inv(−s), "
            "j_ref^dep(−s,K), G_inv, G_dep, G_same, G_avail, G_inv = G_avail + G_dep and Δ_inv(e); Equations (12)–(19) "
            "continue the matrix, composition and selection-stability definitions. The synchronized mapping is supplied in Additional file 4 "
            "under docs/Equation_to_code_mapping.csv. The verified public base remains paper-release-2026-07-r9 at commit "
            "9635a902fa3cc7bb7b71a234c1b2bbbe415193f0; the paper43 ten-seed completion overlay is on public-release hold."
        )
    p = first(doc, "For each held-out seed, one K = 32")
    if p:
        p.text = (
            "For each held-out seed s, j_ref^inv(−s) was selected once from all 32 candidates using the other nine seeds; "
            "j_ref^dep(−s,K) was selected analogously within C_K. For u=(s,f), G_inv and G_dep subtract the utility of ĵ_u(K) "
            "from the respective reference utility, G_same subtracts it from the same-fold eligible maximum, and G_avail is the "
            "K-invariant reference utility minus the K-dependent reference utility. Thus G_inv = G_avail + G_dep exactly. "
            "Δ_inv(e) averages G_inv(u,32) − G_inv(u,4) over the ten seeds and three folds; negative values favour the smaller "
            "full-registry completion gap at K = 32. Practical-equivalence grids were 0.005/0.010/0.020 ROC-AUC and "
            "0.025/0.050/0.100 RMSE, and remain retrospective reporting tolerances rather than minimum important differences."
        )
    doc.save(path)


def mapping_rows():
    impl = "paper43_completion/scripts/analyze_paper43_fixed_reference_and_equivalence_20260726.py"
    return [
        (1, "Validation-selected and finite-audit-best candidates", "src/fzyc_mol/analysis.py and scripts/reproduce_analysis.py", "selected and finite-audit-best candidate identities"),
        (2, "Finite-audit selection loss and range normalization", "src/fzyc_mol/selection/regret_metrics.py", "same-unit selection gap"),
        (3, "CAHit@3 and reciprocal rank", "src/fzyc_mol/selection/ranking_metrics.py", "chance-adjusted Hit@3 inputs and MRR"),
        (4, "Chance MRR expectation and normalization", "src/fzyc_mol/selection/ranking_metrics.py", "normalized MRR gain"),
        (5, "Validation-selected candidate alias j-hat_u(K)", impl, "selected candidate identity"),
        (6, "K-invariant full-registry reference j_ref^inv(-s)", impl, "fixed K=32 cross-fitted reference"),
        (7, "K-dependent within-prefix reference j_ref^dep(-s,K)", impl, "K-dependent cross-fitted reference"),
        (8, "K-invariant and K-dependent gaps G_inv and G_dep", impl, "fixed_k32_gap and k_dependent_crossfit_gap"),
        (9, "Same-fold finite-set opportunity gap G_same", impl, "same_fold_gap"),
        (10, "Availability component and exact decomposition G_inv = G_avail + G_dep", impl, "availability_component and unit-level identity"),
        (11, "Endpoint contrast Delta_inv(e)", impl, "fixed-reference K=32 minus K=4 contrast"),
        (12, "Raw and row-centred utility matrices", "scripts/analyze_paper43_effective_diversity_10seed_20260726.py", "raw and row-centred effective-diversity matrices"),
        (13, "Reference-relative and within-unit-rank matrices", "scripts/analyze_paper43_effective_diversity_10seed_20260726.py", "fixed-reference-relative and rank matrices"),
        (14, "Ledoit-Wolf correlation and eigenvalue proportions", "scripts/analyze_paper43_effective_diversity_10seed_20260726.py", "shrinkage correlations"),
        (15, "Entropy and participation-ratio effective ranks", "scripts/analyze_paper43_effective_diversity_10seed_20260726.py", "effective-rank estimates"),
        (16, "Observed opportunity and selected gain", "paper43 completion source tables", "audit-best opportunity and selected gain"),
        (17, "Composition normalization", "paper43 completion source tables", "normalized composition effects"),
        (18, "Normalized cross-fitted composition gap", impl, "normalized cross-fitted completion gap"),
        (19, "Normalized selection entropy", "src/fzyc_mol/selection/stability_metrics.py", "candidate_selection_entropy_normalized"),
    ]


def write_mapping() -> None:
    destinations = [
        QC / "Equation_to_code_mapping.csv",
        PKG / "working_additional4_r3_20260726" / "docs" / "Equation_to_code_mapping.csv",
        PKG / "repository_deposition" / "paper-release-2026-07-r10_STAGING_CONFLICT_HOLD" / "docs" / "Equation_to_code_mapping.csv",
    ]
    for path in destinations:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(["equation", "estimand_or_definition", "implementation", "manuscript_facing_output", "verification_status"])
            for number, estimand, implementation, output in mapping_rows():
                writer.writerow([number, estimand, implementation, output, "verified against supplied source and synchronized notation"])


def update_workbook() -> None:
    wb = load_workbook(XLSX)
    readme = wb["README"]
    readme["A1"] = "Additional file 2: Machine-readable Supplementary Tables S1–S46"
    for row in range(1, readme.max_row + 1):
        table = readme.cell(row, 1).value
        if table == "Table S6":
            readme.cell(row, 2).value = "S6 Effective diversity"
            readme.cell(row, 3).value = "Ten-seed matrix-dependent effective diversity (outer 30 × K; inner 90 × K)"
        elif table == "Table S7":
            readme.cell(row, 2).value = "S7 Diversity sensitivity"
            readme.cell(row, 3).value = "Leave-one-seed/fold and prespecified-reference sensitivity"
        elif isinstance(readme.cell(row, 3).value, str) and "common to all 15 outer units" in readme.cell(row, 3).value:
            readme.cell(row, 3).value = readme.cell(row, 3).value.replace("all 15 outer units", "all 15 five-seed secondary outer units")
    index = wb["Paper43_Index"]
    index["C2"] = "Δ_inv(e): K-invariant K=32 minus K=4 completion-gap contrasts"
    wb.save(XLSX)


def main() -> None:
    QC.mkdir(parents=True, exist_ok=True)
    shutil.copy2(EN, EN_OLD)
    shutil.copy2(ZH, ZH_OLD)
    shutil.copy2(SUPP, SUPP_OLD)
    revise_english(EN_OLD, EN_REVISED)
    revise_chinese(ZH_OLD, ZH_REVISED)
    revise_supplement(SUPP)
    write_mapping()
    update_workbook()
    print(EN_REVISED)
    print(ZH_REVISED)
    print(SUPP)
    print(XLSX)


if __name__ == "__main__":
    main()
